"""Freeze approved human pilot decisions and build the final review package."""

from __future__ import annotations

import csv
import json
import os
import re
import shutil
from collections import Counter
from collections.abc import Callable
from dataclasses import dataclass
from datetime import UTC, date, datetime
from hashlib import sha256
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from quebradas_limaeste.inventory.full_ingestion import (
    ALL_DOCUMENT_FIELDS,
    ALL_DOCUMENTS_OUTPUT,
)
from quebradas_limaeste.inventory.limaeste_filter import (
    CANDIDATE_FILTER_FIELDS,
    DOCUMENT_FILTER_FIELDS,
    FILTER_CANDIDATES_OUTPUT,
    FILTER_DOCUMENTS_OUTPUT,
)
from quebradas_limaeste.inventory.limaeste_review_package import (
    assign_limaeste_review_filenames,
)

PILOT_VALIDATION_CONFIG = Path("configs/sources/indeci_pilot_human_validation.yaml")
PILOT_HUMAN_DECISIONS_OUTPUT = Path("metadata/indeci/pilot_human_decisions.csv")
PILOT_RETAINED_DOCUMENTS_OUTPUT = Path("metadata/indeci/pilot_retained_documents.csv")
PILOT_EXCLUDED_DOCUMENTS_OUTPUT = Path("metadata/indeci/pilot_excluded_documents.csv")
PILOT_VALIDATION_SUMMARY_OUTPUT = Path("metadata/indeci/pilot_validation_summary.json")
FINAL_PILOT_PACKAGE_OUTPUT = Path("review_packages/cusipata_limaeste_final")

HUMAN_DECISION_FIELDS = (
    "document_id",
    "review_priority",
    "automatic_spatial_relevance",
    "automatic_event_relevance",
    "human_inventory_decision",
    "human_exclusion_reason",
    "human_review_status",
    "reviewed_by",
    "reviewed_at_utc",
    "decision_basis",
)
RETAINED_DOCUMENT_FIELDS = (
    "document_id",
    "priority",
    "year",
    "title",
    "report_type",
    "report_number",
    "report_date",
    "detected_locations",
    "detected_event_terms",
    "rainfall_related",
    "candidate_count",
    "strong_count",
    "moderate_count",
    "weak_count",
    "source_url",
    "raw_local_path",
)
EXCLUDED_DOCUMENT_FIELDS = (
    "document_id",
    "priority",
    "year",
    "title",
    "spatial_relevance",
    "event_relevance",
    "human_exclusion_reason",
    "source_url",
)
FINAL_INDEX_FIELDS = (
    "Prioridad",
    "Ubicación",
    "Evento",
    "Relación con lluvia",
    "Fecha evento",
    "Fecha reporte",
    "Título",
    "Tipo reporte",
    "N.º reporte",
    "Candidate strength máximo",
    "Archivo",
    "Document ID",
    "Decisión humana",
)

_EXCLUSION_REASONS = (
    "outside_study_area",
    "excluded_event_type",
    "outside_area_and_event_type",
    "out_of_scope",
)
_CONTROL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_PRIORITY_ORDER = ("P1", "P2")
_README_TEXT = """FINALIDAD
Paquete final del inventario documental piloto aprobado por el investigador.

ALCANCE
- contiene únicamente documentos P1 y P2 con decisión humana INCLUDE;
- no confirma eventos ni crea ground truth o etiquetas de entrenamiento;
- las clasificaciones automáticas originales permanecen separadas e intactas;
- los documentos PX y sus archivos raw no fueron eliminados.
"""


class PilotValidationError(RuntimeError):
    """Raised when the approved validation freeze cannot be applied safely."""


@dataclass(frozen=True)
class PilotValidationPolicy:
    reviewed_by: str
    decision_basis: str
    approved_priority_decisions: dict[str, str]


@dataclass(frozen=True)
class FrozenDecision:
    document_id: str
    review_priority: str
    automatic_spatial_relevance: str
    automatic_event_relevance: str
    human_inventory_decision: str
    human_exclusion_reason: str


def load_pilot_validation_policy(
    path: Path = PILOT_VALIDATION_CONFIG,
    *,
    allowed_root: Path,
) -> PilotValidationPolicy:
    """Load the versioned human decision approved for this pilot."""
    source = _safe_input(path, root=Path(allowed_root).resolve(), maximum=32_768)
    try:
        payload = json.loads(source.read_text(encoding="utf-8"))
    except (UnicodeError, json.JSONDecodeError) as exc:
        raise PilotValidationError(
            "validation config must be valid UTF-8 JSON"
        ) from exc
    expected = {
        "schema_version",
        "reviewed_by",
        "decision_basis",
        "approved_priority_decisions",
    }
    if not isinstance(payload, dict) or set(payload) != expected:
        raise PilotValidationError("validation config keys do not match schema")
    if payload["schema_version"] != 1:
        raise PilotValidationError("validation config schema version is unsupported")
    reviewed_by = _bounded_text(payload["reviewed_by"], "reviewed_by", maximum=100)
    decision_basis = _bounded_text(
        payload["decision_basis"], "decision_basis", maximum=300
    )
    decisions = payload["approved_priority_decisions"]
    if not isinstance(decisions, dict) or decisions != {
        "P1": "include",
        "P2": "include",
        "PX": "exclude",
    }:
        raise PilotValidationError(
            "approved decisions must be exactly P1/P2 include and PX exclude"
        )
    return PilotValidationPolicy(
        reviewed_by=reviewed_by,
        decision_basis=decision_basis,
        approved_priority_decisions=dict(decisions),
    )


def derive_frozen_decision(
    automatic: dict[str, str],
    *,
    policy: PilotValidationPolicy,
) -> FrozenDecision:
    """Apply only the explicit researcher-approved priority decision."""
    priority = automatic.get("review_priority", "")
    try:
        human_decision = policy.approved_priority_decisions[priority]
    except KeyError as exc:
        raise PilotValidationError(
            f"priority {priority or '<blank>'} has no approved human decision"
        ) from exc
    exclusion_reason = ""
    if human_decision == "exclude":
        outside = automatic.get("spatial_relevance") in {"low", "unknown"}
        excluded_event = automatic.get("event_relevance") == "excluded_topic"
        if outside and excluded_event:
            exclusion_reason = "outside_area_and_event_type"
        elif outside:
            exclusion_reason = "outside_study_area"
        elif excluded_event:
            exclusion_reason = "excluded_event_type"
        else:
            exclusion_reason = "out_of_scope"
    return FrozenDecision(
        document_id=automatic.get("document_id", ""),
        review_priority=priority,
        automatic_spatial_relevance=automatic.get("spatial_relevance", ""),
        automatic_event_relevance=automatic.get("event_relevance", ""),
        human_inventory_decision=human_decision,
        human_exclusion_reason=exclusion_reason,
    )


def execute_pilot_validation_freeze(
    *,
    documents_path: Path = ALL_DOCUMENTS_OUTPUT,
    filtered_documents_path: Path = FILTER_DOCUMENTS_OUTPUT,
    filtered_candidates_path: Path = FILTER_CANDIDATES_OUTPUT,
    config_path: Path = PILOT_VALIDATION_CONFIG,
    decisions_output: Path = PILOT_HUMAN_DECISIONS_OUTPUT,
    retained_output: Path = PILOT_RETAINED_DOCUMENTS_OUTPUT,
    excluded_output: Path = PILOT_EXCLUDED_DOCUMENTS_OUTPUT,
    summary_output: Path = PILOT_VALIDATION_SUMMARY_OUTPUT,
    allowed_root: Path,
    expected_documents: int = 131,
    now: Callable[[], datetime] = lambda: datetime.now(UTC),
) -> dict[str, object]:
    """Freeze the approved human inventory decision without changing A1.16."""
    root = Path(allowed_root).resolve()
    policy = load_pilot_validation_policy(config_path, allowed_root=root)
    documents_source = _safe_input(documents_path, root=root, maximum=20_000_000)
    filtered_source = _safe_input(
        filtered_documents_path, root=root, maximum=20_000_000
    )
    candidates_source = _safe_input(
        filtered_candidates_path, root=root, maximum=20_000_000
    )
    documents, document_fields = _read_csv(documents_source)
    filtered, filtered_fields = _read_csv(filtered_source)
    candidates, candidate_fields = _read_csv(candidates_source)
    if document_fields != ALL_DOCUMENT_FIELDS:
        raise PilotValidationError("all_documents CSV columns do not match schema")
    if filtered_fields != DOCUMENT_FILTER_FIELDS:
        raise PilotValidationError("filtered document CSV columns do not match schema")
    if candidate_fields != CANDIDATE_FILTER_FIELDS:
        raise PilotValidationError("filtered candidate CSV columns do not match schema")
    if len(documents) != expected_documents or len(filtered) != expected_documents:
        raise PilotValidationError("pilot inputs do not contain expected documents")
    document_by_id = _unique_rows(documents, id_field="document_id", label="source")
    filtered_by_id = _unique_rows(filtered, id_field="document_id", label="filtered")
    if set(document_by_id) != set(filtered_by_id):
        raise PilotValidationError("source and filtered document IDs do not match")
    _validate_candidate_rows(candidates, known_documents=set(document_by_id))
    _validate_automatic_consistency(document_by_id, filtered_by_id)
    timestamp = _utc_timestamp(now)

    decisions = []
    retained = []
    excluded = []
    for automatic in sorted(filtered, key=_automatic_sort_key):
        source = document_by_id[automatic["document_id"]]
        decision = derive_frozen_decision(automatic, policy=policy)
        decisions.append(
            {
                "document_id": decision.document_id,
                "review_priority": decision.review_priority,
                "automatic_spatial_relevance": decision.automatic_spatial_relevance,
                "automatic_event_relevance": decision.automatic_event_relevance,
                "human_inventory_decision": decision.human_inventory_decision,
                "human_exclusion_reason": decision.human_exclusion_reason,
                "human_review_status": "reviewed",
                "reviewed_by": policy.reviewed_by,
                "reviewed_at_utc": timestamp,
                "decision_basis": policy.decision_basis,
            }
        )
        if decision.human_inventory_decision == "include":
            retained.append(_retained_row(automatic, source=source))
        else:
            excluded.append(
                _excluded_row(
                    automatic,
                    source=source,
                    reason=decision.human_exclusion_reason,
                )
            )

    retained_ids = {row["document_id"] for row in retained}
    retained_candidates = [
        row for row in candidates if row["document_id"] in retained_ids
    ]
    summary = _validation_summary(
        decisions,
        retained=retained,
        excluded=excluded,
        retained_candidates=retained_candidates,
    )
    destinations = tuple(
        _safe_output(path, root=root)
        for path in (
            decisions_output,
            retained_output,
            excluded_output,
            summary_output,
        )
    )
    if len(set(destinations)) != len(destinations):
        raise PilotValidationError("pilot outputs must use distinct paths")
    if set(destinations) & {documents_source, filtered_source, candidates_source}:
        raise PilotValidationError("pilot outputs must not overwrite automatic inputs")
    if any(path.exists() or path.is_symlink() for path in destinations):
        raise PilotValidationError(
            "pilot validation output already exists; refusing to rewrite "
            "human decisions"
        )
    _write_csv(destinations[0], HUMAN_DECISION_FIELDS, decisions)
    _write_csv(destinations[1], RETAINED_DOCUMENT_FIELDS, retained)
    _write_csv(destinations[2], EXCLUDED_DOCUMENT_FIELDS, excluded)
    _write_json(destinations[3], summary)
    return {
        **summary,
        "decisions_output": str(destinations[0]),
        "retained_output": str(destinations[1]),
        "excluded_output": str(destinations[2]),
        "summary_output": str(destinations[3]),
    }


def build_final_pilot_package(
    *,
    filtered_documents_path: Path = FILTER_DOCUMENTS_OUTPUT,
    decisions_path: Path = PILOT_HUMAN_DECISIONS_OUTPUT,
    retained_documents_path: Path = PILOT_RETAINED_DOCUMENTS_OUTPUT,
    output_dir: Path = FINAL_PILOT_PACKAGE_OUTPUT,
    allowed_root: Path,
    expected_documents: int | None = None,
) -> dict[str, int]:
    """Build a SHA-verified package containing only approved retained PDFs."""
    root = Path(allowed_root).resolve()
    filtered_source = _safe_input(
        filtered_documents_path, root=root, maximum=20_000_000
    )
    decisions_source = _safe_input(decisions_path, root=root, maximum=5_000_000)
    retained_source = _safe_input(retained_documents_path, root=root, maximum=5_000_000)
    filtered, filtered_fields = _read_csv(filtered_source)
    decisions, decision_fields = _read_csv(decisions_source)
    retained, retained_fields = _read_csv(retained_source)
    if filtered_fields != DOCUMENT_FILTER_FIELDS:
        raise PilotValidationError("filtered document CSV columns do not match schema")
    if decision_fields != HUMAN_DECISION_FIELDS:
        raise PilotValidationError("human decision CSV columns do not match schema")
    if retained_fields != RETAINED_DOCUMENT_FIELDS:
        raise PilotValidationError("retained document CSV columns do not match schema")
    if expected_documents is not None and len(retained) != expected_documents:
        raise PilotValidationError(
            f"final package expected {expected_documents} documents; "
            f"found {len(retained)}"
        )
    filtered_by_id = _unique_rows(filtered, id_field="document_id", label="filtered")
    decision_by_id = _unique_rows(decisions, id_field="document_id", label="decision")
    retained_by_id = _unique_rows(retained, id_field="document_id", label="retained")
    included_ids = {
        document_id
        for document_id, row in decision_by_id.items()
        if row["human_inventory_decision"] == "include"
    }
    if set(retained_by_id) != included_ids:
        raise PilotValidationError("retained documents do not match human includes")
    if not included_ids <= set(filtered_by_id):
        raise PilotValidationError("retained document is absent from A1.16 output")
    for document_id in included_ids:
        automatic = filtered_by_id[document_id]
        decision = decision_by_id[document_id]
        if automatic["review_priority"] not in _PRIORITY_ORDER:
            raise PilotValidationError("final package contains a non-P1/P2 document")
        if (
            decision["human_review_status"] != "reviewed"
            or decision["automatic_spatial_relevance"] != automatic["spatial_relevance"]
            or decision["automatic_event_relevance"] != automatic["event_relevance"]
            or decision["review_priority"] != automatic["review_priority"]
        ):
            raise PilotValidationError("human decision does not match automatic source")

    destination = _safe_output(output_dir, root=root)
    if destination.exists() or destination.is_symlink():
        raise PilotValidationError(
            "final package already exists; refusing to overwrite"
        )
    destination.parent.mkdir(parents=True, exist_ok=True)
    staging = destination.with_name(f".{destination.name}.part")
    if staging.exists() or staging.is_symlink():
        raise PilotValidationError("final package staging directory already exists")
    retained_automatic = [filtered_by_id[row["document_id"]] for row in retained]
    filenames = assign_limaeste_review_filenames(retained_automatic)
    copied_names: dict[str, str] = {}
    staging.mkdir()
    try:
        for priority in _PRIORITY_ORDER:
            (staging / priority).mkdir()
        for automatic in sorted(retained_automatic, key=_automatic_sort_key):
            if automatic["local_pdf_status"] == "missing":
                continue
            if automatic["local_pdf_status"] != "available":
                raise PilotValidationError("local PDF status is invalid")
            raw = _safe_raw_pdf(automatic["raw_local_path"], root=root)
            digest = _file_sha256(raw)
            if digest != automatic["sha256"]:
                raise PilotValidationError(
                    f"raw SHA-256 mismatch for {automatic['document_id']}"
                )
            filename = filenames[automatic["document_id"]]
            target = staging / automatic["review_priority"] / filename
            shutil.copyfile(raw, target)
            if _file_sha256(target) != digest:
                raise PilotValidationError(
                    f"copy SHA-256 mismatch for {automatic['document_id']}"
                )
            copied_names[automatic["document_id"]] = filename
        ordered = sorted(retained_automatic, key=_automatic_sort_key)
        index_rows = [
            _final_index_row(row, copied_names=copied_names) for row in ordered
        ]
        _write_package_csv(staging / "INDICE_FINAL.csv", index_rows)
        _write_final_workbook(staging / "INDICE_FINAL.xlsx", rows=index_rows)
        (staging / "README.txt").write_text(
            _README_TEXT,
            encoding="utf-8",
            newline="\n",
        )
        staging.rename(destination)
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    copied_priorities = Counter(
        filtered_by_id[document_id]["review_priority"] for document_id in copied_names
    )
    return {
        "pdfs_p1": copied_priorities["P1"],
        "pdfs_p2": copied_priorities["P2"],
        "excel_rows": len(retained),
        "sha_mismatches": 0,
        "missing_local_pdf": len(retained) - len(copied_names),
    }


def _validation_summary(
    decisions: list[dict[str, object]],
    *,
    retained: list[dict[str, object]],
    excluded: list[dict[str, object]],
    retained_candidates: list[dict[str, str]],
) -> dict[str, object]:
    included_priority = Counter(row["priority"] for row in retained)
    excluded_reasons = Counter(row["human_exclusion_reason"] for row in excluded)
    strengths = Counter(row["candidate_strength"] for row in retained_candidates)
    clusters = {
        row["event_cluster_id"]
        for row in retained_candidates
        if row["event_cluster_id"]
    }
    return {
        "documents_total": len(decisions),
        "documents_included": len(retained),
        "documents_excluded": len(excluded),
        "included_by_priority": {
            priority: included_priority[priority] for priority in _PRIORITY_ORDER
        },
        "excluded_by_reason": {
            reason: excluded_reasons[reason] for reason in _EXCLUSION_REASONS
        },
        "events_in_retained_documents": len(clusters),
        "event_clusters_in_retained_documents": len(clusters),
        "strong_candidates": strengths["strong"],
        "moderate_candidates": strengths["moderate"],
        "weak_candidates": strengths["weak"],
        "human_validation_completed": True,
    }


def _retained_row(
    automatic: dict[str, str],
    *,
    source: dict[str, str],
) -> dict[str, object]:
    return {
        "document_id": automatic["document_id"],
        "priority": automatic["review_priority"],
        "year": automatic["year"],
        "title": automatic["title"],
        "report_type": automatic["report_type"],
        "report_number": automatic["report_number"],
        "report_date": automatic["report_date"],
        "detected_locations": automatic["detected_locations"],
        "detected_event_terms": automatic["detected_event_terms"],
        "rainfall_related": automatic["rainfall_related"],
        "candidate_count": automatic["candidate_count"],
        "strong_count": automatic["strong_count"],
        "moderate_count": automatic["moderate_count"],
        "weak_count": automatic["weak_count"],
        "source_url": source["source_url"],
        "raw_local_path": automatic["raw_local_path"],
    }


def _excluded_row(
    automatic: dict[str, str],
    *,
    source: dict[str, str],
    reason: str,
) -> dict[str, object]:
    return {
        "document_id": automatic["document_id"],
        "priority": automatic["review_priority"],
        "year": automatic["year"],
        "title": automatic["title"],
        "spatial_relevance": automatic["spatial_relevance"],
        "event_relevance": automatic["event_relevance"],
        "human_exclusion_reason": reason,
        "source_url": source["source_url"],
    }


def _final_index_row(
    automatic: dict[str, str],
    *,
    copied_names: dict[str, str],
) -> dict[str, object]:
    strength = next(
        (
            item
            for item in ("strong", "moderate", "weak")
            if int(automatic[f"{item}_count"] or 0) > 0
        ),
        "",
    )
    return {
        "Prioridad": automatic["review_priority"],
        "Ubicación": automatic["primary_location"],
        "Evento": automatic["primary_event"],
        "Relación con lluvia": automatic["rainfall_related"],
        "Fecha evento": _excel_date(automatic["event_date"]),
        "Fecha reporte": _excel_date(automatic["report_date"]),
        "Título": automatic["title"],
        "Tipo reporte": automatic["review_report_type"] or automatic["report_type"],
        "N.º reporte": automatic["review_report_number"] or automatic["report_number"],
        "Candidate strength máximo": strength,
        "Archivo": copied_names.get(automatic["document_id"], ""),
        "Document ID": automatic["document_id"],
        "Decisión humana": "INCLUDE",
    }


def _validate_automatic_consistency(
    documents: dict[str, dict[str, str]],
    filtered: dict[str, dict[str, str]],
) -> None:
    shared = (
        "year",
        "title",
        "report_type",
        "report_number",
        "report_date",
        "relevance_status",
        "matched_terms",
        "raw_local_path",
        "sha256",
    )
    for document_id, automatic in filtered.items():
        source = documents[document_id]
        if any(automatic[field] != source[field] for field in shared):
            raise PilotValidationError(
                f"automatic source fields changed for {document_id}"
            )
        if automatic["review_priority"] not in {"P1", "P2", "PX"}:
            raise PilotValidationError(
                f"priority {automatic['review_priority']} lacks human approval"
            )


def _validate_candidate_rows(
    candidates: list[dict[str, str]],
    *,
    known_documents: set[str],
) -> None:
    if len(candidates) > 5_000:
        raise PilotValidationError("candidate input exceeds row limit")
    candidate_ids = [row["candidate_id"] for row in candidates]
    if len(candidate_ids) != len(set(candidate_ids)):
        raise PilotValidationError("candidate input contains duplicate IDs")
    if any(row["document_id"] not in known_documents for row in candidates):
        raise PilotValidationError("candidate references an unknown document")


def _unique_rows(
    rows: list[dict[str, str]],
    *,
    id_field: str,
    label: str,
) -> dict[str, dict[str, str]]:
    result = {row[id_field]: row for row in rows}
    if "" in result or len(result) != len(rows):
        raise PilotValidationError(f"{label} rows contain blank or duplicate IDs")
    return result


def _automatic_sort_key(row: dict[str, str]) -> tuple[object, ...]:
    priority = row["review_priority"]
    rank = {"P1": 0, "P2": 1, "PX": 2}.get(priority, 3)
    event_date = row["event_date"] or row["report_date"]
    year, month, day = _sortable_date(event_date)
    return (rank, -year, -month, -day, row["document_id"])


def _sortable_date(value: str) -> tuple[int, int, int]:
    match = re.fullmatch(r"(20\d{2})-(\d{2}|XX)-(\d{2}|XX)", value)
    if not match:
        return (0, 0, 0)
    return tuple(int(item) if item != "XX" else 0 for item in match.groups())


def _excel_date(value: str) -> date | str | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError:
        if re.fullmatch(r"20\d{2}-(?:\d{2}|XX)-(?:\d{2}|XX)", value):
            return value
        raise PilotValidationError("pilot review date is not ISO-like") from None


def _write_final_workbook(
    path: Path,
    *,
    rows: list[dict[str, object]],
) -> None:
    workbook = Workbook()
    final = workbook.active
    final.title = "FINAL"
    final.append(FINAL_INDEX_FIELDS)
    for row in rows:
        final.append([_excel_safe(row[field]) for field in FINAL_INDEX_FIELDS])
    final.freeze_panes = "A2"
    final.auto_filter.ref = final.dimensions
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in final[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in final.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.lstrip().startswith(
                ("=", "+", "-", "@")
            ):
                cell.data_type = "s"
                cell.quotePrefix = True
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for field in ("Fecha evento", "Fecha reporte"):
        column = FINAL_INDEX_FIELDS.index(field) + 1
        for row_number in range(2, max(final.max_row, 2) + 1):
            final.cell(row=row_number, column=column).number_format = "yyyy-mm-dd"
    widths = {
        "Prioridad": 10,
        "Ubicación": 25,
        "Evento": 30,
        "Título": 58,
        "Archivo": 68,
        "Document ID": 30,
        "Decisión humana": 18,
    }
    for index, field in enumerate(FINAL_INDEX_FIELDS, start=1):
        final.column_dimensions[get_column_letter(index)].width = widths.get(field, 20)
    final.row_dimensions[1].height = 42

    readme = workbook.create_sheet("README")
    for row_number, line in enumerate(_README_TEXT.splitlines(), start=1):
        readme.cell(row=row_number, column=1, value=line)
    readme["A1"].font = Font(bold=True)
    readme.column_dimensions["A"].width = 90

    summary = workbook.create_sheet("SUMMARY")
    summary.append(("Métrica", "Conteo"))
    priorities = Counter(row["Prioridad"] for row in rows)
    summary.append(("total", len(rows)))
    for priority in _PRIORITY_ORDER:
        summary.append((priority, priorities[priority]))
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 14
    workbook.save(path)


def _read_csv(path: Path) -> tuple[list[dict[str, str]], tuple[str, ...]]:
    try:
        with path.open(newline="", encoding="utf-8") as source:
            reader = csv.DictReader(source)
            fields = tuple(reader.fieldnames or ())
            rows = list(reader)
    except (OSError, UnicodeError, csv.Error) as exc:
        raise PilotValidationError(f"could not read CSV: {path.name}") from exc
    if any(None in row for row in rows):
        raise PilotValidationError(f"CSV has rows wider than its header: {path.name}")
    return rows, fields


def _write_csv(
    path: Path,
    fields: tuple[str, ...],
    rows: list[dict[str, object]],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.part")
    if temporary.exists() or temporary.is_symlink():
        raise PilotValidationError("pilot CSV staging file already exists")
    try:
        with temporary.open("x", newline="", encoding="utf-8") as output:
            writer = csv.DictWriter(
                output,
                fieldnames=fields,
                extrasaction="raise",
                quoting=csv.QUOTE_ALL,
                lineterminator="\n",
            )
            writer.writeheader()
            for row in rows:
                writer.writerow(
                    {field: _spreadsheet_safe(row.get(field, "")) for field in fields}
                )
            output.flush()
            os.fsync(output.fileno())
        temporary.replace(path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def _write_package_csv(path: Path, rows: list[dict[str, object]]) -> None:
    with path.open("x", newline="", encoding="utf-8") as output:
        writer = csv.DictWriter(
            output,
            fieldnames=FINAL_INDEX_FIELDS,
            extrasaction="raise",
            quoting=csv.QUOTE_ALL,
            lineterminator="\n",
        )
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    field: _spreadsheet_safe(row.get(field, ""))
                    for field in FINAL_INDEX_FIELDS
                }
            )
        output.flush()
        os.fsync(output.fileno())


def _write_json(path: Path, payload: dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.part")
    if temporary.exists() or temporary.is_symlink():
        raise PilotValidationError("pilot JSON staging file already exists")
    try:
        with temporary.open("x", encoding="utf-8", newline="\n") as output:
            json.dump(payload, output, ensure_ascii=False, indent=2, sort_keys=True)
            output.write("\n")
            output.flush()
            os.fsync(output.fileno())
        temporary.replace(path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def _spreadsheet_safe(value: object) -> object:
    if not isinstance(value, str):
        return value
    value = _CONTROL_CHARACTERS.sub("", value)
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _excel_safe(value: object) -> object:
    if isinstance(value, str):
        return _CONTROL_CHARACTERS.sub("", value)
    return value


def _utc_timestamp(now: Callable[[], datetime]) -> str:
    value = now()
    if not isinstance(value, datetime) or value.tzinfo is None:
        raise PilotValidationError("review timestamp must be timezone-aware")
    return (
        value.astimezone(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    )


def _bounded_text(value: object, field: str, *, maximum: int) -> str:
    if not isinstance(value, str):
        raise PilotValidationError(f"{field} must be text")
    clean = _CONTROL_CHARACTERS.sub("", value).strip()
    if not clean or len(clean) > maximum:
        raise PilotValidationError(f"{field} has invalid length")
    return clean


def _safe_input(path: Path, *, root: Path, maximum: int) -> Path:
    unresolved = path if path.is_absolute() else root / path
    if unresolved.is_symlink():
        raise PilotValidationError("input symlinks are not allowed")
    try:
        target = unresolved.resolve(strict=True)
    except OSError as exc:
        raise PilotValidationError(f"input does not exist: {path}") from exc
    if not target.is_relative_to(root) or not target.is_file():
        raise PilotValidationError("input is outside the workspace or not a file")
    if target.stat().st_size > maximum:
        raise PilotValidationError("input exceeds its size limit")
    return target


def _safe_output(path: Path, *, root: Path) -> Path:
    unresolved = path if path.is_absolute() else root / path
    if unresolved.is_symlink():
        raise PilotValidationError("output symlinks are not allowed")
    target = unresolved.resolve()
    if not target.is_relative_to(root):
        raise PilotValidationError("output is outside the workspace")
    return target


def _safe_raw_pdf(value: str, *, root: Path) -> Path:
    unresolved = Path(value) if Path(value).is_absolute() else root / value
    if unresolved.is_symlink():
        raise PilotValidationError("raw PDF symlinks are not allowed")
    target = unresolved.resolve()
    raw_root = (root / "data/raw/indeci").resolve()
    if not target.is_relative_to(raw_root) or not target.is_file():
        raise PilotValidationError("raw PDF is outside data/raw/indeci or missing")
    if target.suffix.lower() != ".pdf":
        raise PilotValidationError("raw review source is not a PDF")
    return target


def _file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as source:
        while chunk := source.read(64 * 1024):
            digest.update(chunk)
    return digest.hexdigest()
