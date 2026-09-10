"""Deterministic operational metadata for the INDECI coauthor review package."""

from __future__ import annotations

import csv
import json
import os
import re
import shutil
import unicodedata
from collections import Counter
from collections.abc import Iterable
from datetime import date
from hashlib import sha256
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from quebradas_limaeste.inventory.candidate_audit import (
    AUDIT_FIELDS,
    CONSOLIDATED_FIELDS,
)
from quebradas_limaeste.inventory.document_classification import semantic_text

REVIEW_PACKAGE_OUTPUT = Path("review_packages/indeci_all")
REVIEW_INDEX_FIELDS = (
    "Prioridad",
    "Archivo",
    "Document ID",
    "Año",
    "Título original",
    "Ubicación automática",
    "Evento automático",
    "Fecha evento automática",
    "Fecha reporte",
    "Tipo reporte",
    "N.º reporte",
    "Clasificación automática",
    "Candidate strength máximo",
    "Términos encontrados",
    "Páginas relevantes",
    "N.º candidatos",
    "N.º clusters",
    "golden_control",
    "¿Corresponde a Cusipata?",
    "Evento confirmado",
    "Fecha evento revisada",
    "Incluir en inventario",
    "Precisión espacial",
    "Observación del revisor",
    "Nombre del revisor",
    "Fecha de revisión",
)
HUMAN_REVIEW_FIELDS = (
    "¿Corresponde a Cusipata?",
    "Evento confirmado",
    "Fecha evento revisada",
    "Incluir en inventario",
    "Precisión espacial",
    "Observación del revisor",
    "Nombre del revisor",
    "Fecha de revisión",
)
REVIEW_FILE_MAP_FIELDS = (
    "document_id",
    "raw_local_path",
    "review_filename",
    "review_local_path",
    "sha256",
    "priority",
    "location",
    "event",
    "event_date",
    "report_date",
    "report_type",
    "report_number",
)
_REQUIRED_DOCUMENT_FIELDS = {
    "document_id",
    "year",
    "title",
    "report_type",
    "report_number",
    "report_date",
    "raw_local_path",
    "sha256",
    "page_count",
    "ocr_required",
    "relevance_status",
    "matched_terms",
    "golden_control",
    "review_priority",
    "review_location",
    "review_event",
    "review_event_date",
    "review_report_date",
    "review_report_type",
    "review_report_number",
    "candidate_strength_max",
    "relevant_pages",
    "event_candidate_count",
    "event_cluster_count",
}
_EXPECTED_GOLDENS = {
    "positive_control": "INDECI_RC630_20190303",
    "negative_ambiguous_control": "INDECI_IE1496_20230505",
}
_README_TEXT = """OBJETIVO
Revisar documentos INDECI candidatos para identificar evidencia de eventos en
la quebrada Cusipata / San Bartolomé, distrito de Chaclacayo.

ORDEN
P1 primero
P2 luego
P3 luego
P4 al final

IMPORTANTE
- un PDF relevante para Lima no implica evento en Cusipata;
- verificar siempre ubicación;
- distinguir fecha del reporte y fecha del evento;
- no considerar ausencia de mención como ausencia física del evento;
- no crear etiquetas de entrenamiento;
- registrar incertidumbre como Dudoso/Pendiente.
"""
_WORKBOOK_README = (
    "Revisar PDFs en orden P1 a P4.",
    "No modificar el nombre de archivo.",
    "Completar únicamente las columnas humanas.",
    'Sí en "Incluir en inventario" no significa training_label=1.',
    'Registrar dudas como "Pendiente".',
    "No borrar filas.",
)
_STRENGTH_RANK = {"": -1, "weak": 0, "moderate": 1, "strong": 2}
_REPORT_CODES = {
    "reporte_complementario": "RC",
    "reporte_preliminar": "RP",
    "informe_emergencia": "IE",
}
_EVENT_TYPES = {
    "activacion quebrada": "ACTIVACION-DE-QUEBRADA",
    "activacion de quebrada": "ACTIVACION-DE-QUEBRADA",
    "desborde rio": "DESBORDE",
    "desborde": "DESBORDE",
    "deslizamiento": "DESLIZAMIENTO",
    "erosion fluvial": "EROSION-FLUVIAL",
    "flujo detritos": "FLUJO-DE-DETRITOS",
    "flujo de detritos": "FLUJO-DE-DETRITOS",
    "flujo de lodo": "FLUJO-DE-LODO",
    "huaico": "HUAICO",
    "huayco": "HUAICO",
    "inundacion": "INUNDACION",
    "lluvia intensa": "LLUVIAS-INTENSAS",
    "lluvias intensas": "LLUVIAS-INTENSAS",
}
_EVENT_PREFERENCE = (
    "ACTIVACION-DE-QUEBRADA",
    "FLUJO-DE-DETRITOS",
    "FLUJO-DE-LODO",
    "HUAICO",
    "DESBORDE",
    "INUNDACION",
    "DESLIZAMIENTO",
    "LLUVIAS-INTENSAS",
    "EROSION-FLUVIAL",
)


class ReviewPackageError(RuntimeError):
    """Raised when a portable review package cannot be built safely."""


def build_review_package(
    *,
    documents_path: Path,
    candidates_path: Path,
    clusters_path: Path,
    output_dir: Path = REVIEW_PACKAGE_OUTPUT,
    allowed_root: Path,
    expected_documents: int = 131,
) -> dict[str, int]:
    """Build an immutable-PDF review copy and human-editable index offline."""
    root = Path(allowed_root).resolve()
    documents_source = _safe_input(documents_path, root=root, maximum=20_000_000)
    candidates_source = _safe_input(candidates_path, root=root, maximum=20_000_000)
    clusters_source = _safe_input(clusters_path, root=root, maximum=20_000_000)
    documents, document_fields = _read_csv(documents_source)
    candidates, candidate_fields = _read_csv(candidates_source)
    clusters, cluster_fields = _read_csv(clusters_source)
    if not set(document_fields) >= _REQUIRED_DOCUMENT_FIELDS:
        raise ReviewPackageError("all_documents CSV is missing required fields")
    if "training_label" in document_fields:
        raise ReviewPackageError("all_documents must not contain training_label")
    if candidate_fields != AUDIT_FIELDS or cluster_fields != CONSOLIDATED_FIELDS:
        raise ReviewPackageError("candidate or cluster CSV columns do not match schema")
    if len(documents) != expected_documents:
        raise ReviewPackageError(
            f"review package expected {expected_documents} documents; "
            f"found {len(documents)}"
        )
    document_ids = [row["document_id"] for row in documents]
    if len(set(document_ids)) != len(document_ids):
        raise ReviewPackageError("all_documents contains duplicate document_id values")
    if any(row["validation_status"] != "pending_review" for row in candidates):
        raise ReviewPackageError("candidate decisions must remain pending_review")
    if any(row["validation_status"] != "pending_review" for row in clusters):
        raise ReviewPackageError("cluster decisions must remain pending_review")
    _validate_golden_controls(documents)
    destination = _safe_output(output_dir, root=root)
    if destination.exists() or destination.is_symlink():
        raise ReviewPackageError("review package already exists; refusing to overwrite")
    destination.parent.mkdir(parents=True, exist_ok=True)
    staging = destination.with_name(f".{destination.name}.part")
    if staging.exists() or staging.is_symlink():
        raise ReviewPackageError("review package staging directory already exists")
    staging.mkdir()
    try:
        for priority in ("P1", "P2", "P3", "P4"):
            (staging / priority).mkdir()
        filenames = assign_review_filenames(documents)
        mappings = _copy_review_pdfs(
            documents,
            filenames=filenames,
            staging=staging,
            destination=destination,
            root=root,
        )
        mapped_names = {row["document_id"]: row["review_filename"] for row in mappings}
        index_rows = [_index_row(row, filenames=mapped_names) for row in documents]
        _write_csv(
            staging / "review_file_map.csv",
            REVIEW_FILE_MAP_FIELDS,
            mappings,
        )
        _write_csv(
            staging / "INDICE_REVISION.csv",
            REVIEW_INDEX_FIELDS,
            index_rows,
        )
        (staging / "README.txt").write_text(
            _README_TEXT,
            encoding="utf-8",
            newline="\n",
        )
        _write_workbook(
            staging / "INDICE_REVISION.xlsx",
            rows=index_rows,
            summary=_summary_rows(documents),
        )
        staging.rename(destination)
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    collisions = sum(
        filenames[str(row["document_id"])] != build_review_filename(row)
        for row in documents
    )
    return {
        "pdfs_copied": len(mappings),
        "excel_rows": len(index_rows),
        "filename_collisions": collisions,
        "hash_mismatches": 0,
        "golden_controls_present": len(_EXPECTED_GOLDENS),
    }


def derive_review_fields(
    document: dict[str, object],
    *,
    candidates: Iterable[dict[str, object]] = (),
    clusters: Iterable[dict[str, object]] = (),
) -> dict[str, object]:
    """Derive review ordering labels without making scientific decisions."""
    candidate_rows = tuple(candidates)
    cluster_rows = tuple(clusters)
    location = select_review_location(document, candidates=candidate_rows)
    event, event_date = select_review_event(
        document,
        candidates=candidate_rows,
        clusters=cluster_rows,
    )
    strength = max(
        (str(row.get("candidate_strength") or "") for row in candidate_rows),
        key=lambda value: _STRENGTH_RANK.get(value, -1),
        default="",
    )
    relevance = str(document.get("relevance_status") or "unclassified")
    if strength == "strong" or (
        location in {"CUSIPATA", "SAN-BARTOLOME"} and event != "UNKNOWN"
    ):
        priority = "P1"
    elif strength == "moderate" or (
        location == "CHACLACAYO" and event != "UNKNOWN"
    ):
        priority = "P2"
    elif (
        location not in {"UNKNOWN", "CUSIPATA", "SAN-BARTOLOME", "CHACLACAYO"}
        and event != "UNKNOWN"
        and relevance in {"relevant", "potentially_relevant"}
    ):
        priority = "P3"
    else:
        priority = "P4"
    pages = sorted(
        {
            int(row["source_page"])
            for row in candidate_rows
            if str(row.get("source_page") or "").isdigit()
        }
    )
    report_type = str(document.get("report_type") or "")
    return {
        "review_priority": priority,
        "review_location": location,
        "review_event": event,
        "review_event_date": event_date,
        "review_report_date": str(document.get("report_date") or ""),
        "review_report_type": _REPORT_CODES.get(report_type, "UNKNOWN"),
        "review_report_number": str(document.get("report_number") or ""),
        "candidate_strength_max": strength,
        "relevant_pages": "|".join(map(str, pages)),
        "event_candidate_count": len(candidate_rows),
        "event_cluster_count": len(cluster_rows),
    }


def select_review_location(
    document: dict[str, object],
    *,
    candidates: Iterable[dict[str, object]] = (),
) -> str:
    """Choose the first explicit location in the required preference order."""
    candidate_rows = tuple(candidates)
    matched = _matched_terms(document.get("matched_terms"))
    evidence = " ".join(
        (
            str(document.get("title") or ""),
            *matched,
            *(
                str(row.get(field) or "")
                for row in candidate_rows
                for field in ("reported_quebrada", "site_evidence")
            ),
        )
    )
    normalized = semantic_text(evidence)
    for term, label in (
        ("cusipata", "CUSIPATA"),
        ("san bartolome", "SAN-BARTOLOME"),
        ("chaclacayo", "CHACLACAYO"),
    ):
        if _contains(normalized, term):
            return label
    if _contains(normalized, "lurigancho chosica") or any(
        _contains(normalized, term) for term in ("lurigancho", "chosica")
    ):
        return "LURIGANCHO-CHOSICA"
    for row in candidate_rows:
        reported = str(row.get("reported_quebrada") or "").strip()
        if reported:
            label = sanitize_token(reported, maximum=40)
            if label not in {"", "QUEBRADA", "UNKNOWN"}:
                return label
    title = str(document.get("title") or "")
    explicit = re.search(
        r"\b(?:distrito|provincia)\s+de\s+([A-Za-zÁÉÍÓÚÜÑáéíóúüñ -]{2,50})",
        title,
        flags=re.IGNORECASE,
    )
    if explicit:
        label = sanitize_token(re.split(r"[,;()]", explicit.group(1))[0], maximum=40)
        if label:
            return label
    if _contains(normalized, "lima"):
        return "LIMA"
    return "UNKNOWN"


def select_review_event(
    document: dict[str, object],
    *,
    candidates: Iterable[dict[str, object]] = (),
    clusters: Iterable[dict[str, object]] = (),
) -> tuple[str, str]:
    """Prefer consolidated evidence and never substitute the report date."""
    cluster_rows = tuple(clusters)
    if cluster_rows:
        chosen = min(
            cluster_rows,
            key=lambda row: (
                -_STRENGTH_RANK.get(str(row.get("candidate_strength") or ""), -1),
                str(row.get("event_date") or "9999-99-99"),
                str(row.get("event_cluster_id") or ""),
            ),
        )
        normalized = _normalize_event(chosen.get("event_type"))
        if normalized != "UNKNOWN":
            return normalized, str(chosen.get("event_date") or "")
    candidate_rows = tuple(candidates)
    if candidate_rows:
        chosen = min(
            candidate_rows,
            key=lambda row: (
                -_STRENGTH_RANK.get(str(row.get("candidate_strength") or ""), -1),
                str(row.get("event_date") or "9999-99-99"),
                str(row.get("candidate_id") or ""),
            ),
        )
        normalized = _normalize_event(chosen.get("event_type"))
        if normalized != "UNKNOWN":
            return normalized, str(chosen.get("event_date") or "")
    explicit = {
        _normalize_event(term) for term in _matched_terms(document.get("matched_terms"))
    }
    for event in _EVENT_PREFERENCE:
        if event in explicit:
            return event, ""
    return "UNKNOWN", ""


def assign_review_filenames(
    documents: Iterable[dict[str, object]],
) -> dict[str, str]:
    """Assign deterministic traceable names and resolve collisions by document ID."""
    assigned: dict[str, str] = {}
    used: set[str] = set()
    for document in sorted(documents, key=lambda row: str(row["document_id"])):
        document_id = str(document["document_id"])
        base = build_review_filename(document)
        name = base
        if name.casefold() in used:
            stem = base[:-4]
            suffix = sanitize_token(document_id, maximum=24)[-12:] or "DOCUMENT"
            name = f"{stem}_{suffix}.pdf"
        counter = 2
        while name.casefold() in used:
            name = f"{base[:-4]}_{counter}.pdf"
            counter += 1
        assigned[document_id] = name
        used.add(name.casefold())
    return assigned


def build_review_filename(document: dict[str, object]) -> str:
    """Build one portable review filename from derived, source-linked fields."""
    priority = str(document.get("review_priority") or "P4")
    if priority not in {"P1", "P2", "P3", "P4"}:
        priority = "P4"
    location = sanitize_token(document.get("review_location"), maximum=40) or "UNKNOWN"
    event = sanitize_token(document.get("review_event"), maximum=40) or "UNKNOWN"
    event_date = _date_token(document.get("review_event_date"), prefix="EVT")
    report_date = _date_token(document.get("review_report_date"), prefix="REP")
    report_type = sanitize_token(
        document.get("review_report_type"), maximum=12
    ) or "UNKNOWN"
    report_number = sanitize_token(
        document.get("review_report_number"), maximum=20
    )
    identity = f"{report_type}{report_number}" if report_number else report_type
    return f"{priority}_{location}_{event}_{event_date}_{report_date}_{identity}.pdf"


def sanitize_token(value: object, *, maximum: int = 40) -> str:
    """Return an uppercase ASCII token suitable for portable filenames."""
    text = "" if value is None else str(value)
    decomposed = unicodedata.normalize("NFKD", text)
    ascii_text = "".join(
        character
        for character in decomposed
        if not unicodedata.combining(character) and ord(character) < 128
    )
    token = re.sub(r"[^A-Za-z0-9]+", "-", ascii_text).strip("-").upper()
    return token[:maximum].rstrip("-")


def _normalize_event(value: object) -> str:
    normalized = semantic_text(str(value or ""))
    return _EVENT_TYPES.get(normalized, "UNKNOWN")


def _matched_terms(value: object) -> tuple[str, ...]:
    if isinstance(value, list | tuple):
        return tuple(str(item) for item in value)
    if not isinstance(value, str) or not value:
        return ()
    try:
        parsed = json.loads(value[1:] if value.startswith("'") else value)
    except json.JSONDecodeError:
        return ()
    if not isinstance(parsed, list):
        return ()
    return tuple(str(item) for item in parsed)


def _date_token(value: object, *, prefix: str) -> str:
    text = str(value or "")
    if re.fullmatch(r"20\d{2}-\d{2}-\d{2}", text):
        return f"{prefix}{text}"
    return f"{prefix}-UNKNOWN"


def _contains(text: str, term: str) -> bool:
    return re.search(rf"(?<!\w){re.escape(term)}(?!\w)", text) is not None


def _copy_review_pdfs(
    documents: list[dict[str, str]],
    *,
    filenames: dict[str, str],
    staging: Path,
    destination: Path,
    root: Path,
) -> list[dict[str, object]]:
    mappings = []
    for document in sorted(
        documents,
        key=lambda row: (row["review_priority"], row["document_id"]),
    ):
        raw_value = document["raw_local_path"]
        if not raw_value:
            continue
        raw = _safe_raw_pdf(raw_value, root=root)
        recorded = document["sha256"]
        digest = _file_sha256(raw)
        if digest != recorded:
            raise ReviewPackageError(
                f"raw SHA-256 mismatch for {document['document_id']}"
            )
        priority = document["review_priority"]
        if priority not in {"P1", "P2", "P3", "P4"}:
            raise ReviewPackageError("review priority is invalid")
        filename = filenames[document["document_id"]]
        target = staging / priority / filename
        if target.exists() or target.is_symlink():
            raise ReviewPackageError("review filename collision was not resolved")
        shutil.copyfile(raw, target)
        copy_digest = _file_sha256(target)
        if copy_digest != digest:
            raise ReviewPackageError(
                f"review copy SHA-256 mismatch for {document['document_id']}"
            )
        mappings.append(
            {
                "document_id": document["document_id"],
                "raw_local_path": str(raw),
                "review_filename": filename,
                "review_local_path": str(destination / priority / filename),
                "sha256": digest,
                "priority": priority,
                "location": document["review_location"],
                "event": document["review_event"],
                "event_date": document["review_event_date"],
                "report_date": document["review_report_date"],
                "report_type": document["review_report_type"],
                "report_number": document["review_report_number"],
            }
        )
    return mappings


def _index_row(
    document: dict[str, str],
    *,
    filenames: dict[str, str],
) -> dict[str, object]:
    relevance = {
        "potentially_relevant": "possible",
        "not_relevant": "irrelevant",
    }.get(document["relevance_status"], document["relevance_status"])
    row: dict[str, object] = {
        "Prioridad": document["review_priority"],
        "Archivo": filenames.get(document["document_id"], ""),
        "Document ID": document["document_id"],
        "Año": int(document["year"]),
        "Título original": document["title"],
        "Ubicación automática": document["review_location"],
        "Evento automático": document["review_event"],
        "Fecha evento automática": _optional_date(document["review_event_date"]),
        "Fecha reporte": _optional_date(document["review_report_date"]),
        "Tipo reporte": document["review_report_type"],
        "N.º reporte": document["review_report_number"],
        "Clasificación automática": relevance,
        "Candidate strength máximo": document["candidate_strength_max"],
        "Términos encontrados": _display_terms(document["matched_terms"]),
        "Páginas relevantes": document["relevant_pages"],
        "N.º candidatos": int(document["event_candidate_count"] or 0),
        "N.º clusters": int(document["event_cluster_count"] or 0),
        "golden_control": "yes" if document["golden_control"] else "",
    }
    row.update(dict.fromkeys(HUMAN_REVIEW_FIELDS, ""))
    return row


def _write_workbook(
    path: Path,
    *,
    rows: list[dict[str, object]],
    summary: list[tuple[str, int]],
) -> None:
    workbook = Workbook()
    review = workbook.active
    review.title = "REVIEW"
    review.append(REVIEW_INDEX_FIELDS)
    for row in rows:
        review.append([row[field] for field in REVIEW_INDEX_FIELDS])
    review.freeze_panes = "A2"
    review.auto_filter.ref = review.dimensions
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in review[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in review.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.lstrip().startswith(
                ("=", "+", "-", "@")
            ):
                cell.data_type = "s"
                cell.quotePrefix = True
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    date_fields = (
        "Fecha evento automática",
        "Fecha reporte",
        "Fecha evento revisada",
        "Fecha de revisión",
    )
    for field in date_fields:
        column = REVIEW_INDEX_FIELDS.index(field) + 1
        for row_number in range(2, max(review.max_row, 2) + 1):
            review.cell(row=row_number, column=column).number_format = "yyyy-mm-dd"
    _add_validations(review)
    widths = {
        "Prioridad": 10,
        "Archivo": 68,
        "Document ID": 28,
        "Año": 10,
        "Título original": 58,
        "Ubicación automática": 24,
        "Evento automático": 28,
        "Términos encontrados": 42,
        "Observación del revisor": 48,
        "Nombre del revisor": 24,
    }
    for index, field in enumerate(REVIEW_INDEX_FIELDS, start=1):
        review.column_dimensions[get_column_letter(index)].width = widths.get(field, 20)
    review.row_dimensions[1].height = 42

    readme = workbook.create_sheet("README")
    readme["A1"] = "INSTRUCCIONES"
    readme["A1"].font = Font(bold=True)
    for row_number, instruction in enumerate(_WORKBOOK_README, start=2):
        readme.cell(row=row_number, column=1, value=instruction)
    readme.column_dimensions["A"].width = 85

    summary_sheet = workbook.create_sheet("SUMMARY")
    summary_sheet.append(("Métrica", "Conteo"))
    for metric, count in summary:
        summary_sheet.append((metric, count))
    for cell in summary_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
    summary_sheet.column_dimensions["A"].width = 28
    summary_sheet.column_dimensions["B"].width = 14
    workbook.save(path)


def _add_validations(sheet) -> None:
    options = {
        "¿Corresponde a Cusipata?": '"Sí,No,Dudoso"',
        "Evento confirmado": '"Sí,No,Dudoso"',
        "Incluir en inventario": '"Sí,No,Pendiente"',
        "Precisión espacial": '"A,B,C,D,E,Desconocida"',
    }
    last_row = max(sheet.max_row, 2)
    for field, formula in options.items():
        validation = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
        )
        sheet.add_data_validation(validation)
        column = get_column_letter(REVIEW_INDEX_FIELDS.index(field) + 1)
        validation.add(f"{column}2:{column}{last_row}")


def _summary_rows(documents: list[dict[str, str]]) -> list[tuple[str, int]]:
    priorities = Counter(row["review_priority"] for row in documents)
    relevance = Counter(row["relevance_status"] for row in documents)
    with_candidates = sum(
        int(row["event_candidate_count"] or 0) > 0 for row in documents
    )
    return [
        ("total", len(documents)),
        *((priority, priorities[priority]) for priority in ("P1", "P2", "P3", "P4")),
        ("relevant", relevance["relevant"]),
        ("possible", relevance["potentially_relevant"]),
        ("irrelevant", relevance["not_relevant"]),
        ("OCR required", sum(row["ocr_required"] == "true" for row in documents)),
        ("con candidatos", with_candidates),
        ("sin candidatos", len(documents) - with_candidates),
    ]


def _validate_golden_controls(documents: list[dict[str, str]]) -> None:
    actual = {
        row["golden_control"]: row["document_id"]
        for row in documents
        if row["golden_control"]
    }
    if actual != _EXPECTED_GOLDENS:
        raise ReviewPackageError("required golden controls are missing or changed")


def _read_csv(path: Path) -> tuple[list[dict[str, str]], tuple[str, ...]]:
    try:
        with path.open(newline="", encoding="utf-8") as source:
            reader = csv.DictReader(source)
            fields = tuple(reader.fieldnames or ())
            rows = []
            for row_number, row in enumerate(reader, start=2):
                if None in row:
                    raise ReviewPackageError(
                        f"CSV row {row_number} contains extra columns"
                    )
                rows.append(
                    {field: _restore_cell(row.get(field) or "") for field in fields}
                )
    except (OSError, UnicodeError, csv.Error) as exc:
        raise ReviewPackageError("review package input is not valid UTF-8 CSV") from exc
    return rows, fields


def _write_csv(
    path: Path,
    fields: tuple[str, ...],
    rows: Iterable[dict[str, object]],
) -> None:
    with path.open("x", newline="", encoding="utf-8") as output:
        writer = csv.DictWriter(
            output,
            fieldnames=fields,
            extrasaction="raise",
            quoting=csv.QUOTE_ALL,
            lineterminator="\n",
        )
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {field: _protect_cell(str(row.get(field, ""))) for field in fields}
            )
        output.flush()
        os.fsync(output.fileno())


def _safe_input(path: Path, *, root: Path, maximum: int) -> Path:
    unresolved = path if path.is_absolute() else root / path
    if unresolved.is_symlink():
        raise ReviewPackageError("input symlinks are not allowed")
    target = unresolved.resolve()
    if not target.is_relative_to(root) or not target.is_file():
        raise ReviewPackageError("input path is outside the workspace or missing")
    if target.stat().st_size > maximum:
        raise ReviewPackageError("input exceeds its size limit")
    return target


def _safe_output(path: Path, *, root: Path) -> Path:
    unresolved = path if path.is_absolute() else root / path
    if unresolved.is_symlink():
        raise ReviewPackageError("output symlinks are not allowed")
    target = unresolved.resolve()
    if not target.is_relative_to(root):
        raise ReviewPackageError("output path is outside the workspace")
    return target


def _safe_raw_pdf(value: str, *, root: Path) -> Path:
    unresolved = Path(value) if Path(value).is_absolute() else root / value
    if unresolved.is_symlink():
        raise ReviewPackageError("raw PDF symlinks are not allowed")
    target = unresolved.resolve()
    raw_root = (root / "data/raw/indeci").resolve()
    if raw_root not in target.parents or not target.is_file():
        raise ReviewPackageError("raw PDF is outside data/raw/indeci or missing")
    if target.suffix.lower() != ".pdf":
        raise ReviewPackageError("raw review source is not a PDF")
    return target


def _file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as source:
        while chunk := source.read(64 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _optional_date(value: str) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise ReviewPackageError("review date is not ISO formatted") from exc


def _display_terms(value: str) -> str:
    try:
        terms = json.loads(value)
    except json.JSONDecodeError as exc:
        raise ReviewPackageError("matched_terms is not valid JSON") from exc
    if not isinstance(terms, list) or not all(isinstance(term, str) for term in terms):
        raise ReviewPackageError("matched_terms must be a JSON string list")
    return " | ".join(terms)


def _protect_cell(value: str) -> str:
    return f"'{value}" if value.lstrip().startswith(("=", "+", "-", "@")) else value


def _restore_cell(value: str) -> str:
    if value.startswith("'") and value[1:].lstrip().startswith(("=", "+", "-", "@")):
        return value[1:]
    return value
