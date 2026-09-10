"""Build the reduced offline coauthor package for Cusipata and Lima Este."""

from __future__ import annotations

import csv
import os
import re
import shutil
from collections import Counter
from datetime import date
from hashlib import sha256
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from quebradas_limaeste.inventory.limaeste_filter import (
    DOCUMENT_FILTER_FIELDS,
    FILTER_DOCUMENTS_OUTPUT,
)
from quebradas_limaeste.inventory.review_package import sanitize_token

LIMAESTE_REVIEW_PACKAGE_OUTPUT = Path("review_packages/cusipata_limaeste_filtered")
REVIEW_INDEX_FIELDS = (
    "Prioridad",
    "Ubicación automática",
    "Evento automático",
    "Relación con lluvia",
    "Fecha evento",
    "Fecha reporte",
    "Título original",
    "Tipo reporte",
    "N.º reporte",
    "Candidate strength máximo",
    "Páginas relevantes",
    "Archivo",
    "Document ID",
    "¿Corresponde a Cusipata/entorno?",
    "Evento confirmado",
    "Fecha evento revisada",
    "Incluir en inventario",
    "Precisión espacial",
    "Observación",
    "Revisor",
    "Fecha revisión",
)
HUMAN_REVIEW_FIELDS = (
    "¿Corresponde a Cusipata/entorno?",
    "Evento confirmado",
    "Fecha evento revisada",
    "Incluir en inventario",
    "Precisión espacial",
    "Observación",
    "Revisor",
    "Fecha revisión",
)
_PRIORITIES = ("P1", "P2", "P3", "P4", "PX")
_CONTROL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_REPORT_CODES = {
    "reporte_complementario": "RC",
    "reporte_preliminar": "RP",
    "informe_emergencia": "IE",
}
_README_TEXT = """OBJETIVO
Revisar los documentos INDECI priorizados para Cusipata, Chaclacayo y el
entorno de Lima Este.

ORDEN
Revisar P1, P2, P3 y P4. PX queda fuera de la revisión prioritaria, pero no
representa una conclusión científica ni debe eliminarse.

REGLAS
- verificar ubicación, evento y relación con lluvia en el PDF;
- distinguir la fecha del evento de la fecha del reporte;
- completar únicamente las columnas humanas;
- registrar incertidumbre como Dudoso o Pendiente;
- no crear etiquetas de entrenamiento.
"""
_WORKBOOK_README = (
    "Revisar en orden P1, P2, P3 y P4; PX no es prioritario.",
    "No modificar los nombres de archivo.",
    "Completar únicamente las columnas humanas.",
    "Incluir en inventario no constituye una etiqueta de entrenamiento.",
    "Registrar incertidumbre como Dudoso o Pendiente.",
    "No borrar filas.",
)


class LimaEsteReviewPackageError(RuntimeError):
    """Raised when a filtered review package cannot be built safely."""


def build_limaeste_review_package(
    *,
    documents_path: Path = FILTER_DOCUMENTS_OUTPUT,
    output_dir: Path = LIMAESTE_REVIEW_PACKAGE_OUTPUT,
    allowed_root: Path,
    expected_documents: int = 131,
) -> dict[str, int]:
    """Copy available local PDFs and build blank human-review indexes."""
    root = Path(allowed_root).resolve()
    source = _safe_input(documents_path, root=root, maximum=20_000_000)
    documents, fields = _read_csv(source)
    if fields != DOCUMENT_FILTER_FIELDS:
        raise LimaEsteReviewPackageError(
            "filtered document CSV columns do not match schema"
        )
    if "training_label" in fields:
        raise LimaEsteReviewPackageError("filtered documents contain training_label")
    if len(documents) != expected_documents:
        raise LimaEsteReviewPackageError(
            f"package expected {expected_documents} documents; found {len(documents)}"
        )
    document_ids = [row["document_id"] for row in documents]
    if len(document_ids) != len(set(document_ids)):
        raise LimaEsteReviewPackageError("filtered documents contain duplicate IDs")
    if any(row["review_priority"] not in _PRIORITIES for row in documents):
        raise LimaEsteReviewPackageError("filtered document priority is invalid")

    destination = _safe_output(output_dir, root=root)
    if destination.exists() or destination.is_symlink():
        raise LimaEsteReviewPackageError(
            "review package already exists; refusing to overwrite"
        )
    destination.parent.mkdir(parents=True, exist_ok=True)
    staging = destination.with_name(f".{destination.name}.part")
    if staging.exists() or staging.is_symlink():
        raise LimaEsteReviewPackageError("review package staging directory exists")
    filenames = assign_limaeste_review_filenames(documents)
    base_filenames = {
        row["document_id"]: build_limaeste_review_filename(row) for row in documents
    }
    copied_names: dict[str, str] = {}
    staging.mkdir()
    try:
        for priority in _PRIORITIES:
            (staging / priority).mkdir()
        for document in sorted(documents, key=_document_sort_key):
            if document["local_pdf_status"] == "missing":
                if document["raw_local_path"]:
                    raise LimaEsteReviewPackageError(
                        "missing PDF status conflicts with raw_local_path"
                    )
                continue
            if document["local_pdf_status"] != "available":
                raise LimaEsteReviewPackageError("local PDF status is invalid")
            raw = _safe_raw_pdf(document["raw_local_path"], root=root)
            digest = _file_sha256(raw)
            if digest != document["sha256"]:
                raise LimaEsteReviewPackageError(
                    f"raw SHA-256 mismatch for {document['document_id']}"
                )
            filename = filenames[document["document_id"]]
            target = staging / document["review_priority"] / filename
            if target.exists() or target.is_symlink():
                raise LimaEsteReviewPackageError("filename collision was not resolved")
            shutil.copyfile(raw, target)
            if _file_sha256(target) != digest:
                raise LimaEsteReviewPackageError(
                    f"copy SHA-256 mismatch for {document['document_id']}"
                )
            copied_names[document["document_id"]] = filename

        ordered = sorted(documents, key=_document_sort_key)
        index_rows = [
            _index_row(document, copied_names=copied_names) for document in ordered
        ]
        _write_csv(staging / "INDICE_REVISION.csv", REVIEW_INDEX_FIELDS, index_rows)
        _write_workbook(
            staging / "INDICE_REVISION.xlsx",
            rows=index_rows,
            documents=documents,
        )
        (staging / "README.txt").write_text(
            _README_TEXT,
            encoding="utf-8",
            newline="\n",
        )
        staging.rename(destination)
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    return {
        "copied_pdfs": len(copied_names),
        "excel_rows": len(documents),
        "filename_collisions": sum(
            filenames[document_id] != base_filenames[document_id]
            for document_id in document_ids
        ),
        "hash_mismatches": 0,
        "missing_local_pdf": sum(
            row["local_pdf_status"] == "missing" for row in documents
        ),
    }


def assign_limaeste_review_filenames(
    documents: list[dict[str, str]],
) -> dict[str, str]:
    """Assign deterministic names and disambiguate duplicate metadata."""
    assigned: dict[str, str] = {}
    used: set[str] = set()
    for document in sorted(documents, key=lambda row: row["document_id"]):
        base = build_limaeste_review_filename(document)
        filename = base
        if filename.casefold() in used:
            suffix = sanitize_token(document["document_id"], maximum=24)[-12:]
            filename = f"{base[:-4]}_{suffix or 'DOCUMENT'}.pdf"
        counter = 2
        while filename.casefold() in used:
            filename = f"{base[:-4]}_{counter}.pdf"
            counter += 1
        assigned[document["document_id"]] = filename
        used.add(filename.casefold())
    return assigned


def build_limaeste_review_filename(document: dict[str, str]) -> str:
    """Build a portable filename from filter output without inventing dates."""
    priority = document["review_priority"]
    if priority not in _PRIORITIES:
        raise LimaEsteReviewPackageError("review priority is invalid")
    location = sanitize_token(document["primary_location"], maximum=40) or "UNKNOWN"
    event = sanitize_token(document["primary_event"], maximum=40) or "UNKNOWN"
    event_date = _date_token(document["event_date"], prefix="EVT")
    report_date = _date_token(document["report_date"], prefix="REP")
    report_type = sanitize_token(
        document["review_report_type"], maximum=12
    ) or _REPORT_CODES.get(document["report_type"], "")
    report_number = sanitize_token(document["review_report_number"], maximum=20)
    identity = (
        f"{report_type}{report_number}" if report_type and report_number else "XXXX"
    )
    return f"{priority}_{location}_{event}_{event_date}_{report_date}_{identity}.pdf"


def _index_row(
    document: dict[str, str],
    *,
    copied_names: dict[str, str],
) -> dict[str, object]:
    strength = next(
        (
            name
            for name in ("strong", "moderate", "weak")
            if int(document[f"{name}_count"] or 0) > 0
        ),
        "",
    )
    row: dict[str, object] = {
        "Prioridad": document["review_priority"],
        "Ubicación automática": document["primary_location"],
        "Evento automático": document["primary_event"],
        "Relación con lluvia": document["rainfall_related"],
        "Fecha evento": _excel_date(document["event_date"]),
        "Fecha reporte": _excel_date(document["report_date"]),
        "Título original": document["title"],
        "Tipo reporte": document["review_report_type"] or document["report_type"],
        "N.º reporte": document["review_report_number"] or document["report_number"],
        "Candidate strength máximo": strength,
        "Páginas relevantes": document["relevant_pages"],
        "Archivo": copied_names.get(document["document_id"], ""),
        "Document ID": document["document_id"],
    }
    row.update(dict.fromkeys(HUMAN_REVIEW_FIELDS, ""))
    return row


def _document_sort_key(row: dict[str, str]) -> tuple[object, ...]:
    year, month, day = _sortable_date(row["event_date"] or row["report_date"])
    return (
        _PRIORITIES.index(row["review_priority"]),
        -year,
        -month,
        -day,
        row["document_id"],
    )


def _sortable_date(value: str) -> tuple[int, int, int]:
    match = re.fullmatch(r"(20\d{2})-(\d{2}|XX)-(\d{2}|XX)", value)
    if not match:
        return (0, 0, 0)
    return tuple(int(item) if item != "XX" else 0 for item in match.groups())


def _date_token(value: str, *, prefix: str) -> str:
    if re.fullmatch(r"20\d{2}-(?:\d{2}|XX)-(?:\d{2}|XX)", value):
        return f"{prefix}{value}"
    return f"{prefix}-UNKNOWN"


def _excel_date(value: str) -> date | str | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError:
        if re.fullmatch(r"20\d{2}-(?:\d{2}|XX)-(?:\d{2}|XX)", value):
            return value
        raise LimaEsteReviewPackageError("review date is not ISO-like") from None


def _write_workbook(
    path: Path,
    *,
    rows: list[dict[str, object]],
    documents: list[dict[str, str]],
) -> None:
    workbook = Workbook()
    review = workbook.active
    review.title = "REVIEW"
    review.append(REVIEW_INDEX_FIELDS)
    for row in rows:
        review.append([_excel_safe(row[field]) for field in REVIEW_INDEX_FIELDS])
    review.freeze_panes = "A2"
    review.auto_filter.ref = review.dimensions
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in review[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in review.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.lstrip().startswith(
                ("=", "+", "-", "@")
            ):
                cell.data_type = "s"
                cell.quotePrefix = True
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for field in (
        "Fecha evento",
        "Fecha reporte",
        "Fecha evento revisada",
        "Fecha revisión",
    ):
        column = REVIEW_INDEX_FIELDS.index(field) + 1
        for row_number in range(2, max(review.max_row, 2) + 1):
            review.cell(row=row_number, column=column).number_format = "yyyy-mm-dd"
    _add_validations(review)
    widths = {
        "Prioridad": 10,
        "Ubicación automática": 25,
        "Evento automático": 30,
        "Título original": 58,
        "Archivo": 68,
        "Document ID": 30,
        "Observación": 48,
        "Revisor": 24,
    }
    for index, field in enumerate(REVIEW_INDEX_FIELDS, start=1):
        review.column_dimensions[get_column_letter(index)].width = widths.get(field, 20)
    review.row_dimensions[1].height = 42

    readme = workbook.create_sheet("README")
    readme["A1"] = "INSTRUCCIONES"
    readme["A1"].font = Font(bold=True)
    for row_number, instruction in enumerate(_WORKBOOK_README, start=2):
        readme.cell(row=row_number, column=1, value=instruction)
    readme.column_dimensions["A"].width = 90

    summary = workbook.create_sheet("SUMMARY")
    summary.append(("Métrica", "Conteo"))
    priorities = Counter(row["review_priority"] for row in documents)
    summary.append(("total", len(documents)))
    for priority in _PRIORITIES:
        summary.append((priority, priorities[priority]))
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 14
    workbook.save(path)


def _add_validations(sheet) -> None:
    options = {
        "¿Corresponde a Cusipata/entorno?": '"Sí,No,Dudoso"',
        "Evento confirmado": '"Sí,No,Dudoso"',
        "Incluir en inventario": '"Sí,No,Pendiente"',
        "Precisión espacial": '"A,B,C,D,E,Desconocida"',
    }
    last_row = max(sheet.max_row, 2)
    for field, formula in options.items():
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        sheet.add_data_validation(validation)
        column = get_column_letter(REVIEW_INDEX_FIELDS.index(field) + 1)
        validation.add(f"{column}2:{column}{last_row}")


def _read_csv(path: Path) -> tuple[list[dict[str, str]], tuple[str, ...]]:
    try:
        with path.open(newline="", encoding="utf-8") as source:
            reader = csv.DictReader(source)
            fields = tuple(reader.fieldnames or ())
            rows = []
            for row_number, row in enumerate(reader, start=2):
                if None in row:
                    raise LimaEsteReviewPackageError(
                        f"CSV row {row_number} contains extra columns"
                    )
                rows.append({field: row.get(field) or "" for field in fields})
    except (OSError, UnicodeError, csv.Error) as exc:
        raise LimaEsteReviewPackageError("input is not valid UTF-8 CSV") from exc
    return rows, fields


def _write_csv(
    path: Path,
    fields: tuple[str, ...],
    rows: list[dict[str, object]],
) -> None:
    with path.open("x", newline="", encoding="utf-8") as output:
        writer = csv.DictWriter(
            output,
            fieldnames=fields,
            extrasaction="raise",
            quoting=csv.QUOTE_ALL,
            lineterminator="\n",
        )
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {field: _protect_cell(str(row.get(field, ""))) for field in fields}
            )
        output.flush()
        os.fsync(output.fileno())


def _protect_cell(value: str) -> str:
    value = _CONTROL_CHARACTERS.sub("", value)
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _excel_safe(value: object) -> object:
    if isinstance(value, str):
        return _CONTROL_CHARACTERS.sub("", value)
    return value


def _safe_input(path: Path, *, root: Path, maximum: int) -> Path:
    unresolved = path if path.is_absolute() else root / path
    if unresolved.is_symlink():
        raise LimaEsteReviewPackageError("input symlinks are not allowed")
    target = unresolved.resolve()
    if not target.is_relative_to(root) or not target.is_file():
        raise LimaEsteReviewPackageError("input is outside the workspace or missing")
    if target.stat().st_size > maximum:
        raise LimaEsteReviewPackageError("input exceeds its size limit")
    return target


def _safe_output(path: Path, *, root: Path) -> Path:
    unresolved = path if path.is_absolute() else root / path
    if unresolved.is_symlink():
        raise LimaEsteReviewPackageError("output symlinks are not allowed")
    target = unresolved.resolve()
    if not target.is_relative_to(root):
        raise LimaEsteReviewPackageError("output is outside the workspace")
    return target


def _safe_raw_pdf(value: str, *, root: Path) -> Path:
    unresolved = Path(value) if Path(value).is_absolute() else root / value
    if unresolved.is_symlink():
        raise LimaEsteReviewPackageError("raw PDF symlinks are not allowed")
    target = unresolved.resolve()
    raw_root = (root / "data/raw/indeci").resolve()
    if not target.is_relative_to(raw_root) or not target.is_file():
        raise LimaEsteReviewPackageError(
            "raw PDF is outside data/raw/indeci or missing"
        )
    if target.suffix.lower() != ".pdf":
        raise LimaEsteReviewPackageError("raw review source is not a PDF")
    return target


def _file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as source:
        while chunk := source.read(64 * 1024):
            digest.update(chunk)
    return digest.hexdigest()
