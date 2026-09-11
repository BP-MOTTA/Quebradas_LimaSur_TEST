import csv
import json
from datetime import UTC, datetime
from hashlib import sha256
from pathlib import Path

import openpyxl
import pytest

from quebradas_limaeste.inventory.full_ingestion import ALL_DOCUMENT_FIELDS
from quebradas_limaeste.inventory.limaeste_filter import (
    CANDIDATE_FILTER_FIELDS,
    DOCUMENT_FILTER_FIELDS,
)
from quebradas_limaeste.inventory.pilot_validation import (
    EXCLUDED_DOCUMENT_FIELDS,
    FINAL_INDEX_FIELDS,
    HUMAN_DECISION_FIELDS,
    RETAINED_DOCUMENT_FIELDS,
    PilotValidationError,
    build_final_pilot_package,
    derive_frozen_decision,
    execute_pilot_validation_freeze,
    load_pilot_validation_policy,
)

VALIDATION_CONFIG = Path("configs/sources/indeci_pilot_human_validation.yaml")


def filtered_document(document_id: str, priority: str, **overrides):
    defaults = {
        "P1": ("core", "target", "CHACLACAYO", "HUAICO"),
        "P2": ("near", "target", "QUIRIO", "ACTIVACION-DE-QUEBRADA"),
        "PX": ("low", "target", "LIMA", "LLUVIAS-INTENSAS"),
    }
    spatial, event_relevance, location, event = defaults[priority]
    row = {field: "" for field in DOCUMENT_FILTER_FIELDS}
    row.update(
        {
            "document_id": document_id,
            "year": "2023",
            "title": f"Documento {document_id}",
            "report_type": "informe_emergencia",
            "report_number": "7",
            "relevance_status": "potentially_relevant",
            "spatial_relevance": spatial,
            "event_relevance": event_relevance,
            "rainfall_related": "true",
            "review_priority": priority,
            "priority_reason": "synthetic automatic decision",
            "detected_locations": location,
            "detected_event_terms": event,
            "candidate_count": "0",
            "strong_count": "0",
            "moderate_count": "0",
            "weak_count": "0",
            "review_required": "true",
            "matched_terms": "[]",
            "primary_location": location,
            "primary_event": event,
            "event_date": "2023-03-16",
            "report_date": "2023-03-20",
            "review_report_type": "IE",
            "review_report_number": "7",
            "local_pdf_status": "available",
        }
    )
    row.update(overrides)
    return row


def source_document(document_id: str, raw: Path, digest: str):
    row = {field: "" for field in ALL_DOCUMENT_FIELDS}
    row.update(
        {
            "document_id": document_id,
            "year": "2023",
            "title": f"Documento {document_id}",
            "report_type": "informe_emergencia",
            "report_number": "7",
            "report_date": "2023-03-20",
            "source_url": f"https://example.test/{document_id}",
            "raw_local_path": str(raw),
            "sha256": digest,
            "download_status": "duplicate",
            "extraction_status": "success",
            "ocr_required": "false",
            "relevance_status": "potentially_relevant",
            "matched_terms": "[]",
            "review_required": "true",
        }
    )
    return row


def filtered_candidate(document_id: str):
    row = {field: "" for field in CANDIDATE_FILTER_FIELDS}
    row.update(
        {
            "candidate_id": f"candidate-{document_id}",
            "event_cluster_id": f"cluster-{document_id}",
            "document_id": document_id,
            "event_date": "2023-03-16",
            "event_type": "huaico",
            "reported_quebrada": "Chaclacayo",
            "source_page": "2",
            "evidence_snippet": "Huaico en Chaclacayo.",
            "candidate_strength": "strong",
            "spatial_relevance": "core",
            "event_relevance": "target",
            "rainfall_related": "true",
            "review_priority": "P1",
            "priority_reason": "core location with target event evidence",
        }
    )
    return row


def write_csv(path: Path, fields, rows) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as output:
        writer = csv.DictWriter(output, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def fixture_inputs(tmp_path: Path):
    config = tmp_path / VALIDATION_CONFIG
    config.parent.mkdir(parents=True)
    config.write_bytes(VALIDATION_CONFIG.read_bytes())
    raw_root = tmp_path / "data/raw/indeci/2023"
    raw_root.mkdir(parents=True)
    ids = ("DOC_P1", "DOC_P2", "DOC_OUT", "DOC_EVENT", "DOC_BOTH")
    raw_files = {}
    for document_id in ids:
        raw = raw_root / f"{document_id}.pdf"
        raw.write_bytes(f"%PDF-1.4\n{document_id}\n%%EOF\n".encode())
        raw_files[document_id] = raw
    filtered = [
        filtered_document(
            "DOC_P1",
            "P1",
            candidate_count="1",
            strong_count="1",
            raw_local_path=str(raw_files["DOC_P1"]),
            sha256=sha256(raw_files["DOC_P1"].read_bytes()).hexdigest(),
        ),
        filtered_document(
            "DOC_P2",
            "P2",
            raw_local_path=str(raw_files["DOC_P2"]),
            sha256=sha256(raw_files["DOC_P2"].read_bytes()).hexdigest(),
        ),
        filtered_document(
            "DOC_OUT",
            "PX",
            raw_local_path=str(raw_files["DOC_OUT"]),
            sha256=sha256(raw_files["DOC_OUT"].read_bytes()).hexdigest(),
        ),
        filtered_document(
            "DOC_EVENT",
            "PX",
            spatial_relevance="core",
            event_relevance="excluded_topic",
            primary_location="CHACLACAYO",
            primary_event="INCENDIO-URBANO",
            detected_locations="CHACLACAYO",
            detected_event_terms="INCENDIO-URBANO",
            raw_local_path=str(raw_files["DOC_EVENT"]),
            sha256=sha256(raw_files["DOC_EVENT"].read_bytes()).hexdigest(),
        ),
        filtered_document(
            "DOC_BOTH",
            "PX",
            event_relevance="excluded_topic",
            primary_event="INCENDIO-FORESTAL",
            detected_event_terms="INCENDIO-FORESTAL",
            raw_local_path=str(raw_files["DOC_BOTH"]),
            sha256=sha256(raw_files["DOC_BOTH"].read_bytes()).hexdigest(),
        ),
    ]
    sources = [
        source_document(
            row["document_id"],
            raw_files[row["document_id"]],
            row["sha256"],
        )
        for row in filtered
    ]
    metadata = tmp_path / "metadata/indeci"
    documents_path = metadata / "all_documents.csv"
    filtered_path = metadata / "geographic_event_filter_documents.csv"
    candidates_path = metadata / "geographic_event_filter_candidates.csv"
    write_csv(documents_path, ALL_DOCUMENT_FIELDS, sources)
    write_csv(filtered_path, DOCUMENT_FILTER_FIELDS, filtered)
    write_csv(
        candidates_path,
        CANDIDATE_FILTER_FIELDS,
        [filtered_candidate("DOC_P1")],
    )
    return config, documents_path, filtered_path, candidates_path, raw_files


def test_freeze_records_explicit_human_decisions_without_changing_automatic_fields(
    tmp_path,
) -> None:
    config, documents, filtered, candidates, raw_files = fixture_inputs(tmp_path)
    raw_before = {
        path: sha256(path.read_bytes()).hexdigest() for path in raw_files.values()
    }
    automatic_before = {
        path: sha256(path.read_bytes()).hexdigest()
        for path in (documents, filtered, candidates)
    }
    reviewed_at = datetime(2026, 9, 10, 18, 30, tzinfo=UTC)

    payload = execute_pilot_validation_freeze(
        documents_path=documents,
        filtered_documents_path=filtered,
        filtered_candidates_path=candidates,
        config_path=config.relative_to(tmp_path),
        allowed_root=tmp_path,
        expected_documents=5,
        now=lambda: reviewed_at,
    )

    assert payload["documents_total"] == 5
    assert payload["documents_included"] == 2
    assert payload["documents_excluded"] == 3
    metadata = tmp_path / "metadata/indeci"
    with (metadata / "pilot_human_decisions.csv").open(
        newline="", encoding="utf-8"
    ) as source:
        decisions = list(csv.DictReader(source))
    with (metadata / "pilot_retained_documents.csv").open(
        newline="", encoding="utf-8"
    ) as source:
        retained = list(csv.DictReader(source))
    with (metadata / "pilot_excluded_documents.csv").open(
        newline="", encoding="utf-8"
    ) as source:
        excluded = list(csv.DictReader(source))
    assert tuple(decisions[0]) == HUMAN_DECISION_FIELDS
    assert tuple(retained[0]) == RETAINED_DOCUMENT_FIELDS
    assert tuple(excluded[0]) == EXCLUDED_DOCUMENT_FIELDS
    by_id = {row["document_id"]: row for row in decisions}
    assert by_id["DOC_P1"]["human_inventory_decision"] == "include"
    assert by_id["DOC_P2"]["human_inventory_decision"] == "include"
    assert by_id["DOC_OUT"]["human_inventory_decision"] == "exclude"
    assert by_id["DOC_OUT"]["human_exclusion_reason"] == "outside_study_area"
    assert by_id["DOC_EVENT"]["human_exclusion_reason"] == "excluded_event_type"
    assert by_id["DOC_BOTH"]["human_exclusion_reason"] == (
        "outside_area_and_event_type"
    )
    assert by_id["DOC_P1"]["automatic_spatial_relevance"] == "core"
    assert by_id["DOC_P1"]["automatic_event_relevance"] == "target"
    assert by_id["DOC_P1"]["review_priority"] == "P1"
    assert all(row["human_review_status"] == "reviewed" for row in decisions)
    assert all(row["reviewed_by"] == "pilot_investigator" for row in decisions)
    assert all(row["reviewed_at_utc"] == "2026-09-10T18:30:00Z" for row in decisions)
    assert "training_label" not in decisions[0]
    assert {row["document_id"] for row in retained} == {"DOC_P1", "DOC_P2"}
    assert {row["priority"] for row in retained} == {"P1", "P2"}
    assert {row["document_id"] for row in excluded} == {
        "DOC_OUT",
        "DOC_EVENT",
        "DOC_BOTH",
    }
    summary = json.loads(
        (metadata / "pilot_validation_summary.json").read_text(encoding="utf-8")
    )
    assert summary["documents_total"] == len(decisions)
    assert summary["documents_included"] == len(retained)
    assert summary["documents_excluded"] == len(excluded)
    assert summary["included_by_priority"] == {"P1": 1, "P2": 1}
    assert summary["excluded_by_reason"] == {
        "excluded_event_type": 1,
        "out_of_scope": 0,
        "outside_area_and_event_type": 1,
        "outside_study_area": 1,
    }
    assert summary["events_in_retained_documents"] == 1
    assert summary["strong_candidates"] == 1
    assert summary["moderate_candidates"] == 0
    assert summary["weak_candidates"] == 0
    assert summary["human_validation_completed"] is True
    assert all(path.exists() for path in raw_files.values())
    assert raw_before == {
        path: sha256(path.read_bytes()).hexdigest() for path in raw_files.values()
    }
    assert automatic_before == {
        path: sha256(path.read_bytes()).hexdigest()
        for path in (documents, filtered, candidates)
    }


def test_frozen_decision_is_reproducible_and_not_a_training_label() -> None:
    policy = load_pilot_validation_policy(
        VALIDATION_CONFIG,
        allowed_root=Path.cwd(),
    )
    automatic = filtered_document("DOC_P1", "P1")

    first = derive_frozen_decision(automatic, policy=policy)
    second = derive_frozen_decision(dict(automatic), policy=policy)

    assert first == second
    assert first.human_inventory_decision == "include"
    assert not hasattr(first, "training_label")


def test_priority_without_explicit_human_approval_is_rejected() -> None:
    policy = load_pilot_validation_policy(
        VALIDATION_CONFIG,
        allowed_root=Path.cwd(),
    )
    automatic = filtered_document("DOC_P3", "P1", review_priority="P3")

    with pytest.raises(PilotValidationError, match="no approved human decision"):
        derive_frozen_decision(automatic, policy=policy)


def test_final_package_contains_only_retained_p1_and_p2(tmp_path) -> None:
    config, documents, filtered, candidates, raw_files = fixture_inputs(tmp_path)
    execute_pilot_validation_freeze(
        documents_path=documents,
        filtered_documents_path=filtered,
        filtered_candidates_path=candidates,
        config_path=config.relative_to(tmp_path),
        allowed_root=tmp_path,
        expected_documents=5,
        now=lambda: datetime(2026, 9, 10, tzinfo=UTC),
    )
    prior_package = tmp_path / "review_packages/cusipata_limaeste_filtered"
    prior_package.mkdir(parents=True)
    marker = prior_package / "immutable-marker.txt"
    marker.write_text("unchanged", encoding="utf-8")
    raw_before = {
        path: sha256(path.read_bytes()).hexdigest() for path in raw_files.values()
    }

    payload = build_final_pilot_package(
        filtered_documents_path=filtered,
        decisions_path=Path("metadata/indeci/pilot_human_decisions.csv"),
        retained_documents_path=Path("metadata/indeci/pilot_retained_documents.csv"),
        output_dir=Path("review_packages/cusipata_limaeste_final"),
        allowed_root=tmp_path,
        expected_documents=2,
    )

    package = tmp_path / "review_packages/cusipata_limaeste_final"
    assert payload == {
        "pdfs_p1": 1,
        "pdfs_p2": 1,
        "excel_rows": 2,
        "sha_mismatches": 0,
        "missing_local_pdf": 0,
    }
    assert {path.name for path in package.iterdir() if path.is_dir()} == {"P1", "P2"}
    assert not (package / "PX").exists()
    assert marker.read_text(encoding="utf-8") == "unchanged"
    copied = list(package.rglob("*.pdf"))
    assert len(copied) == 2
    assert {path.parent.name for path in copied} == {"P1", "P2"}
    assert {sha256(path.read_bytes()).hexdigest() for path in copied} == {
        raw_before[raw_files["DOC_P1"]],
        raw_before[raw_files["DOC_P2"]],
    }
    with (package / "INDICE_FINAL.csv").open(newline="", encoding="utf-8") as source:
        index = list(csv.DictReader(source))
    assert tuple(index[0]) == FINAL_INDEX_FIELDS
    assert [row["Prioridad"] for row in index] == ["P1", "P2"]
    assert all(row["Decisión humana"] == "INCLUDE" for row in index)
    assert "training_label" not in index[0]
    workbook = openpyxl.load_workbook(package / "INDICE_FINAL.xlsx")
    assert workbook.sheetnames == ["FINAL", "README", "SUMMARY"]
    assert workbook["FINAL"].max_row == 3
    assert "training_label" not in [cell.value for cell in workbook["FINAL"][1]]
    assert all(path.exists() for path in raw_files.values())
    assert raw_before == {
        path: sha256(path.read_bytes()).hexdigest() for path in raw_files.values()
    }
