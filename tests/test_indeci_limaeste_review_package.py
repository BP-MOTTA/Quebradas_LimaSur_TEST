import csv
import json
from hashlib import sha256
from pathlib import Path

import openpyxl

from quebradas_limaeste.inventory.limaeste_filter import DOCUMENT_FILTER_FIELDS
from quebradas_limaeste.inventory.limaeste_review_package import (
    HUMAN_REVIEW_FIELDS,
    REVIEW_INDEX_FIELDS,
    assign_limaeste_review_filenames,
    build_limaeste_review_package,
)


def filtered_document(document_id: str, **overrides) -> dict[str, object]:
    row = {field: "" for field in DOCUMENT_FILTER_FIELDS}
    row.update(
        {
            "document_id": document_id,
            "year": "2023",
            "title": "Lluvias intensas en Chaclacayo",
            "report_type": "informe_emergencia",
            "report_number": "7",
            "relevance_status": "potentially_relevant",
            "spatial_relevance": "core",
            "event_relevance": "target",
            "rainfall_related": "true",
            "review_priority": "P1",
            "priority_reason": "core location with target event evidence",
            "detected_locations": "CHACLACAYO",
            "detected_event_terms": "LLUVIAS-INTENSAS",
            "candidate_count": "1",
            "strong_count": "1",
            "moderate_count": "0",
            "weak_count": "0",
            "review_required": "true",
            "matched_terms": '["Chaclacayo","lluvias intensas"]',
            "primary_location": "CHACLACAYO",
            "primary_event": "LLUVIAS-INTENSAS",
            "event_date": "2023-03-16",
            "report_date": "2023-03-20",
            "review_report_type": "IE",
            "review_report_number": "7",
            "relevant_pages": "2",
            "local_pdf_status": "available",
        }
    )
    row.update(overrides)
    return row


def write_csv(path: Path, rows) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as output:
        writer = csv.DictWriter(output, fieldnames=DOCUMENT_FILTER_FIELDS)
        writer.writeheader()
        writer.writerows(rows)


def test_filtered_review_filenames_resolve_metadata_collisions() -> None:
    first = filtered_document("INDECI_IE7_A")
    second = filtered_document("INDECI_IE7_B")

    names = assign_limaeste_review_filenames([second, first])

    assert names[first["document_id"]] != names[second["document_id"]]
    assert names[first["document_id"]].endswith("_IE7.pdf")
    assert names[second["document_id"]].endswith("_INDECI-IE7-B.pdf")


def test_filtered_review_package_is_offline_traceable_and_review_safe(
    tmp_path,
) -> None:
    raw_root = tmp_path / "data/raw/indeci/2023"
    first_raw = raw_root / "first.pdf"
    second_raw = raw_root / "second.pdf"
    raw_root.mkdir(parents=True)
    first_raw.write_bytes(b"%PDF-1.4\nfirst synthetic PDF\n%%EOF\n")
    second_raw.write_bytes(b"%PDF-1.4\nsecond synthetic PDF\n%%EOF\n")
    first_hash = sha256(first_raw.read_bytes()).hexdigest()
    second_hash = sha256(second_raw.read_bytes()).hexdigest()
    rows = [
        filtered_document(
            "INDECI_IE7_A",
            raw_local_path=str(first_raw),
            sha256=first_hash,
            event_date="2023-03-15",
        ),
        filtered_document(
            "INDECI_IE7_B",
            title="=2+2\x01",
            raw_local_path=str(second_raw),
            sha256=second_hash,
            event_date="2023-03-16",
        ),
        filtered_document(
            "INDECI_IE8",
            title="Incendio forestal en Lima",
            spatial_relevance="low",
            event_relevance="excluded_topic",
            rainfall_related="unknown",
            review_priority="PX",
            priority_reason="principal topic is explicitly excluded",
            detected_locations="LIMA",
            detected_event_terms="INCENDIO-FORESTAL",
            candidate_count="0",
            strong_count="0",
            primary_location="LIMA",
            primary_event="INCENDIO-FORESTAL",
            event_date="",
            report_number="8",
            review_report_number="8",
            raw_local_path="",
            sha256="",
            local_pdf_status="missing",
        ),
    ]
    source = tmp_path / "metadata/indeci/geographic_event_filter_documents.csv"
    write_csv(source, rows)
    before = {
        path: (path.name, sha256(path.read_bytes()).hexdigest())
        for path in (first_raw, second_raw)
    }

    result = build_limaeste_review_package(
        documents_path=source,
        output_dir=Path("review_packages/cusipata_limaeste_filtered"),
        allowed_root=tmp_path,
        expected_documents=3,
    )

    package = tmp_path / "review_packages/cusipata_limaeste_filtered"
    assert result == {
        "copied_pdfs": 2,
        "excel_rows": 3,
        "filename_collisions": 0,
        "hash_mismatches": 0,
        "missing_local_pdf": 1,
    }
    assert {path.name for path in package.iterdir() if path.is_dir()} == {
        "P1",
        "P2",
        "P3",
        "P4",
        "PX",
    }
    copied = sorted(package.rglob("*.pdf"))
    assert len(copied) == 2
    assert all(path.parent.name == "P1" for path in copied)
    assert copied[0].name != copied[1].name
    assert all("EVT2023-03-" in path.name for path in copied)
    assert all("REP2023-03-20_IE7" in path.name for path in copied)
    assert all(
        (path.name, sha256(path.read_bytes()).hexdigest()) == original
        for path, original in before.items()
    )
    assert {sha256(path.read_bytes()).hexdigest() for path in copied} == {
        first_hash,
        second_hash,
    }

    with (package / "INDICE_REVISION.csv").open(
        newline="", encoding="utf-8"
    ) as source_file:
        index = list(csv.DictReader(source_file))
    assert tuple(index[0]) == REVIEW_INDEX_FIELDS
    assert [row["Document ID"] for row in index] == [
        "INDECI_IE7_B",
        "INDECI_IE7_A",
        "INDECI_IE8",
    ]
    assert all(row[field] == "" for row in index for field in HUMAN_REVIEW_FIELDS)
    assert index[-1]["Archivo"] == ""
    assert index[0]["Título original"] == "'=2+2"
    assert "training_label" not in index[0]

    workbook = openpyxl.load_workbook(package / "INDICE_REVISION.xlsx")
    assert workbook.sheetnames == ["REVIEW", "README", "SUMMARY"]
    review = workbook["REVIEW"]
    assert tuple(cell.value for cell in review[1]) == REVIEW_INDEX_FIELDS
    assert review.freeze_panes == "A2"
    assert review.auto_filter.ref == review.dimensions
    assert len(review.data_validations.dataValidation) == 4
    title_cell = review.cell(
        row=2,
        column=REVIEW_INDEX_FIELDS.index("Título original") + 1,
    )
    assert title_cell.value == "=2+2"
    assert title_cell.data_type == "s"
    assert all(
        review.cell(row=row, column=REVIEW_INDEX_FIELDS.index(field) + 1).value is None
        for row in range(2, review.max_row + 1)
        for field in HUMAN_REVIEW_FIELDS
    )
    assert "training_label" not in json.dumps(
        [[cell.value for cell in row] for row in review.iter_rows()],
        ensure_ascii=False,
        default=str,
    )
