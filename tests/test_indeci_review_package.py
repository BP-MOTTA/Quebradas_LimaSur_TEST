import csv
import json
from hashlib import sha256
from pathlib import Path

import openpyxl
import pymupdf

from quebradas_limaeste.inventory.candidate_audit import (
    AUDIT_FIELDS,
    CONSOLIDATED_FIELDS,
)
from quebradas_limaeste.inventory.full_ingestion import ALL_DOCUMENT_FIELDS
from quebradas_limaeste.inventory.review_package import (
    HUMAN_REVIEW_FIELDS,
    REVIEW_INDEX_FIELDS,
    assign_review_filenames,
    build_review_filename,
    build_review_package,
    derive_review_fields,
    select_review_event,
)


def base_document(document_id: str, **overrides) -> dict[str, object]:
    row = {field: "" for field in ALL_DOCUMENT_FIELDS}
    row.update(
        {
            "document_id": document_id,
            "discovery_id": "indeci-discovery-" + "a" * 20,
            "year": "2023",
            "title": "Lluvias intensas en Lima",
            "report_type": "informe_emergencia",
            "report_number": "1496",
            "report_date": "2023-05-05",
            "source_connector": "seed_discovery",
            "source_url": "https://portal.indeci.gob.pe/informe/1496/",
            "pdf_url": "https://portal.indeci.gob.pe/raw/1496.pdf",
            "file_size": "1",
            "page_count": "1",
            "download_status": "duplicate",
            "extraction_status": "success",
            "ocr_required": "false",
            "relevance_status": "potentially_relevant",
            "site_match": "false",
            "geographic_match": "true",
            "event_match": "true",
            "matched_terms": '["Lima","lluvias intensas"]',
            "review_required": "true",
            "review_priority": "P3",
            "review_location": "LIMA",
            "review_event": "LLUVIAS-INTENSAS",
            "review_event_date": "",
            "review_report_date": "2023-05-05",
            "review_report_type": "IE",
            "review_report_number": "1496",
            "candidate_strength_max": "",
            "relevant_pages": "",
            "event_candidate_count": "0",
            "event_cluster_count": "0",
        }
    )
    row.update(overrides)
    return row


def candidate(document_id: str, *, strength: str = "strong") -> dict[str, object]:
    row = {field: "" for field in AUDIT_FIELDS}
    row.update(
        {
            "candidate_id": f"candidate-{document_id.casefold().replace('_', '-')}",
            "source_document_id": document_id,
            "source_page": "2",
            "event_date": "2019-02-25",
            "evidence_snippet": "Se reporto huaico en Cusipata.",
            "matched_terms": '["Cusipata","huaico"]',
            "validation_status": "pending_review",
            "reported_quebrada": "Cusipata",
            "canonical_site_id": "quebrada_cusipata_chaclacayo",
            "event_type": "huaico",
            "site_evidence": "Cusipata",
            "event_evidence": "huaico",
            "date_evidence": "25 de febrero de 2019",
            "candidate_strength": strength,
            "review_reason": "explicit_evidence",
        }
    )
    return row


def cluster(document_id: str, *, strength: str = "strong") -> dict[str, object]:
    item = candidate(document_id, strength=strength)
    row = {field: "" for field in CONSOLIDATED_FIELDS}
    row.update(
        {
            "event_cluster_id": f"cluster-{document_id.casefold().replace('_', '-')}",
            "canonical_site_id": "quebrada_cusipata_chaclacayo",
            "event_date": "2019-02-25",
            "event_type": "huaico",
            "reported_quebrada": "Cusipata",
            "candidate_strength": strength,
            "supporting_candidates": item["candidate_id"],
            "supporting_pages": "2",
            "best_evidence_snippet": "Se reporto huaico en Cusipata.",
            "validation_status": "pending_review",
        }
    )
    return row


def write_csv(path: Path, fields, rows) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as output:
        writer = csv.DictWriter(output, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def pdf_bytes(text: str) -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((50, 50), text)
    payload = document.tobytes()
    document.close()
    return payload


def test_review_fields_obey_location_event_date_and_priority_rules() -> None:
    strong = candidate("INDECI_RC630_20190303")
    consolidated = cluster("INDECI_RC630_20190303")
    p1 = derive_review_fields(
        base_document(
            "INDECI_RC630_20190303",
            title="Huaico en Chaclacayo",
            report_date="2019-03-03",
        ),
        candidates=[strong],
        clusters=[consolidated],
    )
    p2 = derive_review_fields(
        base_document(
            "INDECI_IE1_20230101",
            title="Lluvias intensas en Chaclacayo",
            matched_terms='["Chaclacayo","lluvias intensas"]',
        )
    )
    p3 = derive_review_fields(base_document("INDECI_IE2_20230101"))
    p4 = derive_review_fields(
        base_document(
            "INDECI_IE3_20230101",
            title="Incendio urbano",
            relevance_status="not_relevant",
            matched_terms="[]",
        )
    )

    assert p1["review_location"] == "CUSIPATA"
    assert p1["review_event"] == "HUAICO"
    assert p1["review_event_date"] == "2019-02-25"
    assert p1["review_report_date"] == "2019-03-03"
    priorities = [
        p1["review_priority"],
        p2["review_priority"],
        p3["review_priority"],
        p4["review_priority"],
    ]
    assert priorities == [
        "P1",
        "P2",
        "P3",
        "P4",
    ]


def test_report_date_is_never_used_as_missing_event_date() -> None:
    fields = derive_review_fields(
        base_document(
            "INDECI_IE4_20230101",
            title="Lluvias intensas en Lima",
            report_date="2023-01-01",
        )
    )

    assert fields["review_event"] == "LLUVIAS-INTENSAS"
    assert fields["review_event_date"] == ""
    assert "EVT-UNKNOWN" in build_review_filename(fields)
    assert "REP2023-01-01" in build_review_filename(fields)


def test_consolidated_event_types_are_normalized_for_review() -> None:
    cases = {
        "activacion_quebrada": "ACTIVACION-DE-QUEBRADA",
        "flujo_detritos": "FLUJO-DE-DETRITOS",
        "desborde_rio": "DESBORDE",
        "lluvias_intensas": "LLUVIAS-INTENSAS",
    }

    for source, expected in cases.items():
        event, event_date = select_review_event(
            {},
            clusters=[
                {
                    "event_cluster_id": source,
                    "event_type": source,
                    "event_date": "2023-03-16",
                    "candidate_strength": "weak",
                }
            ],
        )
        assert event == expected
        assert event_date == "2023-03-16"


def test_review_filenames_are_sanitized_deterministic_and_collision_safe() -> None:
    first = base_document(
        "INDECI_RC7_20190401",
        review_priority="P1",
        review_location="San Bartolomé / quebrada",
        review_event="flujo de detritos",
        review_event_date="2019-04-01",
        review_report_date="",
        review_report_type="RC",
        review_report_number="7",
    )
    second = {**first, "document_id": "INDECI_RC7_20190401_ALT"}

    one = assign_review_filenames([first, second])
    two = assign_review_filenames([second, first])

    assert one == two
    assert one[first["document_id"]] == (
        "P1_SAN-BARTOLOME-QUEBRADA_FLUJO-DE-DETRITOS_"
        "EVT2019-04-01_REP-UNKNOWN_RC7.pdf"
    )
    assert one[second["document_id"]].endswith("_20190401-ALT.pdf")
    assert one[first["document_id"]] != one[second["document_id"]]


def test_build_review_package_copies_pdfs_and_creates_safe_excel(tmp_path) -> None:
    raw_root = tmp_path / "data/raw/indeci"
    positive_raw = raw_root / "2019" / "ORIGINAL RC630.pdf"
    negative_raw = raw_root / "2023" / "ORIGINAL IE1496.pdf"
    positive_raw.parent.mkdir(parents=True)
    negative_raw.parent.mkdir(parents=True)
    positive_raw.write_bytes(pdf_bytes("positive control"))
    negative_raw.write_bytes(pdf_bytes("negative ambiguous control"))
    positive = base_document(
        "INDECI_RC630_20190303",
        discovery_id="indeci-discovery-" + "b" * 20,
        year="2019",
        title="Huaico en Cusipata y San Bartolome",
        report_type="reporte_complementario",
        report_number="630",
        report_date="2019-03-03",
        raw_local_path=str(positive_raw),
        original_filename=positive_raw.name,
        sha256=sha256(positive_raw.read_bytes()).hexdigest(),
        file_size=str(positive_raw.stat().st_size),
        relevance_status="relevant",
        matched_terms='["Cusipata","Chaclacayo","huaico"]',
        golden_control="positive_control",
        review_priority="P1",
        review_location="CUSIPATA",
        review_event="HUAICO",
        review_event_date="2019-02-25",
        review_report_date="2019-03-03",
        review_report_type="RC",
        review_report_number="630",
        candidate_strength_max="strong",
        relevant_pages="2",
        event_candidate_count="1",
        event_cluster_count="1",
    )
    negative = base_document(
        "INDECI_IE1496_20230505",
        raw_local_path=str(negative_raw),
        original_filename=negative_raw.name,
        sha256=sha256(negative_raw.read_bytes()).hexdigest(),
        file_size=str(negative_raw.stat().st_size),
        golden_control="negative_ambiguous_control",
        review_priority="P4",
    )
    documents = tmp_path / "metadata/indeci/all_documents.csv"
    candidates = tmp_path / "metadata/indeci/all_event_candidates.csv"
    clusters = tmp_path / "metadata/indeci/all_event_clusters.csv"
    write_csv(documents, ALL_DOCUMENT_FIELDS, [positive, negative])
    write_csv(candidates, AUDIT_FIELDS, [candidate(positive["document_id"])])
    write_csv(clusters, CONSOLIDATED_FIELDS, [cluster(positive["document_id"])])
    original_names = {path: path.name for path in (positive_raw, negative_raw)}

    result = build_review_package(
        documents_path=documents,
        candidates_path=candidates,
        clusters_path=clusters,
        output_dir=Path("review_packages/indeci_all"),
        allowed_root=tmp_path,
        expected_documents=2,
    )

    package = tmp_path / "review_packages/indeci_all"
    assert result == {
        "pdfs_copied": 2,
        "excel_rows": 2,
        "filename_collisions": 0,
        "hash_mismatches": 0,
        "golden_controls_present": 2,
    }
    assert {path.name for path in package.rglob("*.pdf")} == {
        "P1_CUSIPATA_HUAICO_EVT2019-02-25_REP2019-03-03_RC630.pdf",
        "P4_LIMA_LLUVIAS-INTENSAS_EVT-UNKNOWN_REP2023-05-05_IE1496.pdf",
    }
    assert all(path.name == name for path, name in original_names.items())
    with (package / "review_file_map.csv").open(newline="", encoding="utf-8") as source:
        mappings = list(csv.DictReader(source))
    assert len(mappings) == 2
    for mapping in mappings:
        raw_digest = sha256(Path(mapping["raw_local_path"]).read_bytes()).hexdigest()
        copy_digest = sha256(
            Path(mapping["review_local_path"]).read_bytes()
        ).hexdigest()
        assert raw_digest == mapping["sha256"]
        assert copy_digest == mapping["sha256"]

    workbook = openpyxl.load_workbook(package / "INDICE_REVISION.xlsx")
    assert workbook.sheetnames == ["REVIEW", "README", "SUMMARY"]
    review = workbook["REVIEW"]
    headers = [cell.value for cell in review[1]]
    assert tuple(headers) == REVIEW_INDEX_FIELDS
    assert "golden_control" in headers
    assert "training_label" not in headers
    assert review.freeze_panes == "A2"
    assert review.auto_filter.ref == review.dimensions
    assert len(review.data_validations.dataValidation) == 4
    human_columns = [headers.index(field) + 1 for field in HUMAN_REVIEW_FIELDS]
    assert all(
        review.cell(row=row, column=column).value is None
        for row in range(2, review.max_row + 1)
        for column in human_columns
    )
    golden_values = {
        review.cell(row=row, column=headers.index("golden_control") + 1).value
        for row in (2, 3)
    }
    assert golden_values == {"yes"}
    assert "training_label" not in json.dumps(
        [[cell.value for cell in row] for row in review.iter_rows()],
        ensure_ascii=False,
        default=str,
    )
